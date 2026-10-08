"""Testes diretos das funcoes que os testes de tela nao exercitam: PDF, Excel da rota, KML e upload.
Importa o app em pasta temporaria (nao toca no seu .db, nas fotos nem no secrets.toml).
uso (na pasta do projeto):  python tests/gircp_exportacoes.py"""
import importlib.util, io, os, re, shutil, sys, tempfile

APP = "app_relatorio.py"
ok = [True]

def check(nome, cond, extra=""):
    ok[0] &= bool(cond)
    print(("PASSOU  " if cond else "FALHOU  ") + nome + (f"  [{extra}]" if extra else ""))

def carregar():
    pasta = tempfile.mkdtemp(prefix="gircp_exp_")
    src = open(APP, encoding="utf-8").read()
    src = re.sub(r'^DB_NAME = ".*"$', f'DB_NAME = r"{pasta}/t.db"', src, flags=re.M)
    src = re.sub(r'^FOTOS_DIR = ".*"$', f'FOTOS_DIR = r"{pasta}/fotos"', src, flags=re.M)
    destino = os.path.join(pasta, "app_teste.py")
    open(destino, "w", encoding="utf-8").write(src)
    os.chdir(pasta)
    spec = importlib.util.spec_from_file_location("gircp_app", destino)
    m = importlib.util.module_from_spec(spec); sys.modules["gircp_app"] = m; spec.loader.exec_module(m)
    return m, pasta

def texto_pdf(b):
    from pypdf import PdfReader
    t = " ".join(p.extract_text() for p in PdfReader(io.BytesIO(b)).pages)
    return re.sub(r"\s+", " ", t)

def main():
    cwd = os.getcwd()
    m, pasta = carregar()
    try:
        IM, MO = "Imediato (0–24h)", "Monitorar"
        dados = {"titulo": "Laudo T", "contato": "Ana", "empresa": "EMP", "telefone": "000", "site_id": "SMXX1", "endereco": "Rua X",
                 "data_hora": "08/09/2026 13:53", "tecnico": "Tecnico A", "art_rrt": "CRT - 1", "numero_relatorio": "GIRCP-1",
                 "revisao": "0", "conclusao": "ok", "status_laudo": "Aprovado"}
        mat = lambda d, q, c: {"descricao": d, "unidade": "un", "quantidade": q, "custo_unit": c}
        foto = lambda t, sev, p, ms: {"titulo": t, "severidade": sev, "prazo_correcao": p, "categoria": "Geral", "comentarios": "c", "materiais": ms, "type": "image/jpeg"}

        # --- funcoes unicas de leitura e custo
        check("_carregar_lista_json: None, vazio, lixo e objeto viram []", all(m._carregar_lista_json(v) == [] for v in (None, "", "lixo", "{}", "null")))
        check("_carregar_lista_json: lista valida e lista ja pronta", m._carregar_lista_json('[1, 2]') == [1, 2] and m._carregar_lista_json([3]) == [3])
        check("_custo_item: numerico, texto e nulo", m._custo_item(mat("a", 2, 3.5)) == 7.0 and m._custo_item(mat("a", "dez", 4.0)) == 4.0 and m._custo_item(None) == 0.0)

        # --- PDF do laudo
        f1 = [foto("A", "Crítico", IM, [mat("cabo", 10, 5.0), mat("borne", 2, 0)]), foto("B", "Normal", MO, [mat("poste", 1, 1500.0), mat("disj", 3, 40.5)])]
        b, _ = m.gerar_pdf(dados, f1, [])
        t = texto_pdf(b)
        check("PDF do laudo e gerado (%PDF)", b[:4] == b"%PDF")
        check("PDF: custo estimado = soma correta (R$ 1671.50)", re.search(r"CUSTO ESTIMADO\s*R\$\s*1671\.50", t) is not None)
        b2, _ = m.gerar_pdf(dados, [foto("A", "Normal", MO, [mat("fita", 2.5, 3.0)])], [])
        qtd = re.findall(r"fita\s+un\s+(\S+)\s", texto_pdf(b2))
        check("PDF: quantidade decimal aparece inteira (2.5, nao 2)", qtd[:1] == ["2.5"], f"mostrou {qtd[:1]}")
        try:
            b3, _ = m.gerar_pdf(dados, [foto("A", "Normal", MO, [mat("cabo", "dez", 5.0)])], [])
            check("PDF: quantidade escrita como texto nao derruba a geracao", b3[:4] == b"%PDF")
        except Exception as e:
            check("PDF: quantidade escrita como texto nao derruba a geracao", False, f"{type(e).__name__}: {str(e)[:50]}")

        # --- PDF e Excel da rota
        rota = [{"id": "BASE", "lon": -46.81, "lat": -23.50}, {"id": "SMA01", "lon": -46.80, "lat": -23.51}, {"id": "SMB02", "lon": -46.79, "lat": -23.52}]
        evs = {"SMA01": [foto("Poste", "Crítico", IM, [mat("poste", 1, 100.0)])], "SMB02": []}
        pb, _ = m.gerar_pdf_rota(rota, 12.3, 1800, "Tecnico A", evs)
        check("PDF da rota e gerado e cita os sites", pb[:4] == b"%PDF" and "SMA01" in texto_pdf(pb))
        xb, _ = m.gerar_excel_rota(rota, 12.3, 1800, "Tecnico A", evs)
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(xb))
        celulas = " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value is not None)
        check("Excel da rota abre e traz os sites", len(wb.sheetnames) >= 1 and "SMA01" in celulas and "SMB02" in celulas, f"abas={wb.sheetnames}")

        # --- KML
        kml = ('<?xml version="1.0" encoding="UTF-8"?><kml xmlns="http://www.opengis.net/kml/2.2"><Document>'
               '<Placemark><name>SMAAA1</name><Point><coordinates>-46.80,-23.51,0</coordinates></Point></Placemark>'
               '<Placemark><name>SMBBB2</name><ExtendedData><Data name="ENDEREÇO"><value>Rua Y, 10</value></Data>'
               '<Data name="LATITUDE"><value>-23,52</value></Data><Data name="LONGITUDE"><value>-46,79</value></Data></ExtendedData></Placemark>'
               '</Document></kml>')
        df = m._parse_kml_to_dataframe(io.BytesIO(kml.encode("utf-8")))
        check("KML: 2 placemarks com colunas esperadas", list(df.columns) == ["SITE", "ENDEREÇO", "GRUPO", "LATITUDE", "LONGITUDE"] and len(df) == 2)
        check("KML: coordenadas de <coordinates> e de ExtendedData (virgula decimal)", abs(df.loc[0, "LATITUDE"] + 23.51) < 1e-9 and abs(df.loc[1, "LONGITUDE"] + 46.79) < 1e-9 and df.loc[1, "ENDEREÇO"] == "Rua Y, 10")

        # --- upload e imagem
        from PIL import Image
        buf = io.BytesIO(); Image.new("RGB", (40, 30), "red").save(buf, format="JPEG"); jpg = buf.getvalue()
        check("upload: JPEG valido e aceito", m._validar_upload("a.jpg", jpg) is None)
        check("upload: vazio, texto disfarcado e grande demais sao recusados",
              "vazio" in (m._validar_upload("a.jpg", b"") or "") and "não é imagem" in (m._validar_upload("a.jpg", b"texto qualquer " * 20) or "")
              and "excede" in (m._validar_upload("a.jpg", b"x" * (m._UPLOAD_MAX_BYTES + 1)) or ""))
        check("comprimir_para_pdf devolve JPEG", m.comprimir_para_pdf(jpg)[:2] == b"\xff\xd8")
    finally:
        os.chdir(cwd); shutil.rmtree(pasta, ignore_errors=True)
    print("\nRESULTADO:", "TUDO PASSOU" if ok[0] else "HA FALHAS")
    return 0 if ok[0] else 1

if __name__ == "__main__":
    sys.exit(main())
