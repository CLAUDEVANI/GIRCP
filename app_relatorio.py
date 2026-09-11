"""
GIRCP — Gerador Inteligente de Relatórios e Controle Fotográfico
| v3.7.1 (Sprints 1, 2 e 3 + Correção SQLite Row)
"""

import streamlit as st
import sqlite3
import base64
import json
import os
import io
import hashlib
import hmac
import math
import requests
import html as html_mod
import pandas as pd
import plotly.express as px
import pydeck as pdk
import xml.etree.ElementTree as ET
import filetype
import qrcode
from datetime import datetime, timedelta
from PIL import Image
from weasyprint import HTML
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ==============================================================================
# 0. CONTROLE DE ACESSO E SEGURANÇA (LGPD)
# ==============================================================================
def check_password():
    if "tentativas" not in st.session_state:
        st.session_state["tentativas"] = 0
        
    if "bloqueado_ate" in st.session_state:
        if datetime.now() < st.session_state["bloqueado_ate"]:
            restante = (st.session_state["bloqueado_ate"] - datetime.now()).seconds
            st.error(f"🔒 Acesso bloqueado. Tente novamente em {restante} segundos.")
            return False
        else:
            del st.session_state["bloqueado_ate"]
            st.session_state["tentativas"] = 0

    def password_entered():
        senha_correta = st.secrets.get("senha_acesso", "GIRCP2026")
        if hmac.compare_digest(st.session_state["password_input"].encode(), senha_correta.encode()):
            st.session_state["password_correct"] = True
            del st.session_state["password_input"]
            st.session_state["tentativas"] = 0
        else:
            st.session_state["password_correct"] = False
            st.session_state["tentativas"] += 1
            if st.session_state["tentativas"] >= 5:
                st.session_state["bloqueado_ate"] = datetime.now() + timedelta(minutes=5)

    if st.session_state.get("password_correct", False):
        return True

    st.markdown("### 🔒 GIRCP - Acesso Restrito")
    st.text_input("Digite a senha de acesso", type="password", on_change=password_entered, key="password_input")
    
    if "password_correct" in st.session_state and not st.session_state["password_correct"]:
        st.error(f"😕 Senha incorreta. Tentativas: {st.session_state['tentativas']}/5")
    return False

# ══════════════════════════════════════════════════════════════════════════
# CONSTANTES E CONFIGURAÇÕES
# ══════════════════════════════════════════════════════════════════════════
COR_AZUL        = "#002060"
COR_AZUL_MED    = "#003087"
COR_AZUL_LIGHT  = "#EBF0FA"
COR_VERMELHO    = "#DA291C"
COR_CINZA       = "#64748B"
COR_CINZA_LIGHT = "#F1F5F9"
COR_BORDA       = "#CBD5E1"
COR_TEXTO       = "#1E293B"
COR_VERDE       = "#16A34A"
COR_AMARELO     = "#D97706"

DB_NAME = "laudos_corp_v3.db"
FOTOS_DIR = "banco_fotos_gircp"
LBL_TITULO = "TÍTULO"
LBL_DESCRICAO = "DESCRIÇÃO"
OPT_DIGITAR_MANUAL = "-- Digitar Manualmente --"
OPT_SELECIONE = "-- SELECIONE --"

os.makedirs(FOTOS_DIR, exist_ok=True)

def init_db():
    with sqlite3.connect(DB_NAME) as conn:
        c = conn.cursor()
        c.execute("PRAGMA journal_mode=WAL")
        c.execute('''
            CREATE TABLE IF NOT EXISTS relatorios (
                id          INTEGER PRIMARY KEY AUTOINCREMENT,
                titulo      TEXT,
                contato     TEXT,
                empresa     TEXT,
                telefone    TEXT,
                email       TEXT,
                site_id     TEXT NOT NULL,
                endereco    TEXT,
                data_hora   TEXT,
                fotos_json  TEXT DEFAULT '[]',
                extras_json TEXT DEFAULT '[]',
                criado_em   TEXT DEFAULT (datetime('now','localtime'))
            )
        ''')
        
        colunas_novas = [
            "criado_em TEXT", "endereco TEXT", "tecnico TEXT", "numero_relatorio TEXT",
            "revisao TEXT", "latitude REAL", "longitude REAL", "art_rrt TEXT",
            "conclusao TEXT", "status_laudo TEXT", "hash_integridade TEXT"
        ]
        for col in colunas_novas:
            try:
                c.execute(f"ALTER TABLE relatorios ADD COLUMN {col}")
            except sqlite3.OperationalError:
                pass

        c.execute("CREATE TABLE IF NOT EXISTS seq_relatorio (ultimo INTEGER DEFAULT 0)")
        c.execute("INSERT INTO seq_relatorio SELECT 0 WHERE NOT EXISTS (SELECT 1 FROM seq_relatorio)")
        
        c.execute('''CREATE TABLE IF NOT EXISTS audit_log (
            id          INTEGER PRIMARY KEY AUTOINCREMENT,
            criado_em   TEXT DEFAULT (datetime('now','localtime')),
            tecnico     TEXT,
            acao        TEXT,
            id_relatorio INTEGER,
            detalhe     TEXT
        )''')
        conn.commit()

def registrar_auditoria(acao: str, id_relatorio: int = None, detalhe: str = ""):
    tecnico = st.session_state.get("_tecnico_global", "sistema")
    with sqlite3.connect(DB_NAME) as conn:
        conn.execute(
            "INSERT INTO audit_log (tecnico, acao, id_relatorio, detalhe) VALUES (?,?,?,?)",
            (tecnico, acao, id_relatorio, detalhe)
        )

def sanitizar(texto: str) -> str:
    return html_mod.escape(str(texto or '').strip())

def comprimir_para_pdf(raw: bytes, max_px: int = 1200, qualidade: int = 78) -> bytes:
    img = Image.open(io.BytesIO(raw)).convert('RGB')
    img.thumbnail((max_px, max_px), Image.Resampling.LANCZOS)
    buf = io.BytesIO()
    img.save(buf, format='JPEG', quality=qualidade, optimize=True)
    return buf.getvalue()

def proximo_numero_relatorio(conn) -> str:
    conn.execute("UPDATE seq_relatorio SET ultimo = ultimo + 1")
    ultimo = conn.execute("SELECT ultimo FROM seq_relatorio").fetchone()[0]
    return f"GIRCP-{datetime.now().year}-{ultimo:04d}"

def _carregar_b64(caminho: str) -> str:
    if os.path.exists(caminho):
        with open(caminho, "rb") as f:
            return base64.b64encode(f.read()).decode('utf-8')
    return ""

def _gerar_qrcode_b64(texto: str) -> str:
    qr = qrcode.make(texto)
    buf = io.BytesIO()
    qr.save(buf, format="PNG")
    return base64.b64encode(buf.getvalue()).decode()

def _botao_backup_db():
    buf = io.BytesIO()
    with open(DB_NAME, "rb") as f:
        buf.write(f.read())
    nome = f"GIRCP_backup_{datetime.now().strftime('%Y%m%d_%H%M%S')}.db"
    st.sidebar.download_button("💾 Backup do Banco", buf.getvalue(), nome, "application/octet-stream")

def _css_pdf() -> str:
    return f"""
    @page {{ size: A4; margin: 14mm 14mm 18mm 14mm; 
        @bottom-left {{ content: "CONFIDENCIAL • USO INTERNO • Dados protegidos pela LGPD | Sistema Corporativo GIRCP"; font-size: 7pt; color: {COR_CINZA}; font-family: 'Segoe UI', Arial, sans-serif; }}
        @bottom-right {{ content: "Página " counter(page) " de " counter(pages) " | HASH_PLACEHOLDER"; font-size: 7pt; color: {COR_CINZA}; font-family: 'Segoe UI', Arial, sans-serif; }}
    }}
    * {{ box-sizing: border-box; margin: 0; padding: 0; }}
    body {{ font-family: 'Segoe UI', Helvetica, Arial, sans-serif; color: {COR_TEXTO}; font-size: 9.5pt; line-height: 1.5; background: #fff; }}
    .watermark {{ position: fixed; top: 38%; left: 50%; transform: translate(-50%, -50%); z-index: -999; width: 65%; opacity: 0.06; }}
    .watermark img {{ width: 100%; }}
    .page-header {{ display: flex; justify-content: space-between; align-items: flex-end; border-bottom: 3px solid {COR_AZUL}; padding-bottom: 10px; margin-bottom: 22px; }}
    .header-titulo {{ color: {COR_AZUL}; font-size: 19pt; font-weight: 900; text-transform: uppercase; letter-spacing: 0.8px; }}
    .header-acento {{ display: inline-block; width: 36px; height: 4px; background: {COR_VERMELHO}; margin-bottom: 4px; }}
    .header-data {{ text-align: right; color: {COR_CINZA}; font-size: 9pt; }}
    .header-data strong {{ color: {COR_AZUL}; font-size: 11pt; }}
    .section-header {{ background: {COR_AZUL}; color: #fff; font-weight: 700; font-size: 10pt; padding: 7px 14px; margin-top: 22px; margin-bottom: 14px; text-transform: uppercase; border-left: 5px solid {COR_VERMELHO}; border-radius: 2px; }}
    .dados-table {{ width: 100%; border-collapse: collapse; margin-bottom: 18px; border: 1px solid {COR_BORDA}; border-radius: 4px; overflow: hidden; }}
    .dados-table td {{ border: 1px solid {COR_BORDA}; padding: 9px 13px; vertical-align: middle; }}
    .dados-table td.label {{ background: {COR_AZUL_LIGHT}; font-weight: 700; color: {COR_AZUL}; text-transform: uppercase; font-size: 7.8pt; width: 20%; }}
    .dados-table td.value {{ color: {COR_TEXTO}; font-size: 9.5pt; width: 30%; }}
    .card-evidencia {{ width: 100%; border-collapse: collapse; margin-bottom: 14px; border: 1px solid {COR_BORDA}; border-radius: 6px; overflow: hidden; page-break-inside: avoid; background: #fff; }}
    .card-evidencia td {{ vertical-align: top; }}
    .card-num {{ background: {COR_AZUL}; color: #fff; font-size: 8pt; font-weight: 700; padding: 3px 8px; text-align: center; writing-mode: vertical-rl; min-width: 22px; }}
    .col-foto {{ width: 43%; background: {COR_CINZA_LIGHT}; padding: 10px; text-align: center; border-right: 1px solid {COR_BORDA}; }}
    .col-foto img {{ width: 100%; max-height: 210px; object-fit: contain; border-radius: 3px; background: #fff; }}
    .col-texto {{ width: 57%; padding: 14px 16px; }}
    .foto-titulo {{ font-weight: 800; color: {COR_AZUL}; font-size: 11pt; text-transform: uppercase; border-bottom: 2px solid {COR_AZUL_LIGHT}; padding-bottom: 7px; margin-bottom: 10px; }}
    .label-desc {{ font-size: 8pt; font-weight: 700; color: {COR_AZUL}; margin-bottom: 4px; }}
    .foto-desc {{ font-size: 9pt; color: {COR_TEXTO}; background: {COR_CINZA_LIGHT}; padding: 9px 11px; border-left: 3px solid {COR_VERMELHO}; border-radius: 0 4px 4px 0; line-height: 1.5; white-space: pre-wrap; }}
    .assinatura-wrapper {{ margin: 44px auto 0 auto; width: 300px; text-align: center; page-break-inside: avoid; }}
    .assinatura-img {{ max-width: 180px; height: auto; display: block; margin: 0 auto -10px auto; }}
    .assinatura-linha {{ border-top: 1.5px solid {COR_TEXTO}; margin: 0 0 6px 0; }}
    .assinatura-nome {{ font-weight: 700; color: {COR_AZUL}; font-size: 9.5pt; }}
    .assinatura-cargo {{ font-size: 8pt; color: {COR_CINZA}; }}
    .logo-wrapper {{ margin-top: 14px; }}
    .logo-img {{ max-width: 140px; height: auto; display: block; margin: 0 auto; }}
    .badge {{ display:inline-block; padding:2px 10px; border-radius:12px; font-size:7.5pt; font-weight:700; margin-bottom:6px; }}
    .badge-critico  {{ background:#DA291C; color:#fff; }}
    .badge-obs      {{ background:#D97706; color:#fff; }}
    .badge-normal   {{ background:#16A34A; color:#fff; }}
    .badge-prazo-imediato {{ background: #7f1d1d; color: #fff; }}
    .badge-prazo-urgente {{ background: #ea580c; color: #fff; }}
    .badge-prazo-planejado {{ background: #0284c7; color: #fff; }}
    .badge-prazo-monitorar {{ background: #475569; color: #fff; }}
    """

def _obter_b64_de_foto(f: dict) -> str:
    if "caminho" in f and os.path.exists(f["caminho"]):
        with open(f["caminho"], "rb") as img_file:
            return base64.b64encode(img_file.read()).decode('utf-8')
    return f.get('base64', '')

def gerar_pdf(dados: dict, fotos: list, extras: list = None) -> tuple[bytes, str]:
    extras = extras or []
    b64_wm  = _carregar_b64("logo_empresa.png")
    b64_sig = _carregar_b64("assinatura_tecnico.png")
    b64_logo = _carregar_b64("logo_empresa.png")

    wm_html = f'<div class="watermark"><img src="data:image/jpeg;base64,{b64_wm}"/></div>' if b64_wm else ""
    sig_img = f'<img class="assinatura-img" src="data:image/png;base64,{b64_sig}"/>' if b64_sig else '<div style="height:40px;"></div>'
    logo_img = f'<img class="logo-img" src="data:image/png;base64,{b64_logo}"/>' if b64_logo else ""

    n_criticos = sum(1 for f in fotos if f.get('severidade','') in ('Critico','Crítico'))
    n_obs      = sum(1 for f in fotos if f.get('severidade','') in ('Observacao','Observação'))
    n_normal   = len(fotos) - n_criticos - n_obs
    all_mats   = [m for f in fotos for m in (f.get('materiais') or []) if m.get('descricao','').strip()]
    total_custo = sum(m.get('quantidade',0)*m.get('custo_unit',0) for m in all_mats)

    resumo_html = f"""
    <div class="section-header">2 &nbsp; RESUMO EXECUTIVO</div>
    <table class="dados-table">
      <tr>
        <td class="label">TOTAL DE EVIDÊNCIAS</td><td class="value">{len(fotos)}</td>
        <td class="label">ITENS CRÍTICOS</td><td class="value" style="color:#DA291C;font-weight:bold;">{n_criticos}</td>
      </tr>
      <tr>
        <td class="label">OBSERVAÇÕES</td><td class="value">{n_obs}</td>
        <td class="label">NORMAIS</td><td class="value">{n_normal}</td>
      </tr>
      <tr>
        <td class="label">TOTAL DE MATERIAIS</td><td class="value">{len(all_mats)} itens</td>
        <td class="label">CUSTO ESTIMADO</td><td class="value" style="font-weight:bold;">R$ {total_custo:.2f}</td>
      </tr>
    </table>
    """

    fotos_html = ""
    for i, f in enumerate(fotos, 1):
        b64 = _obter_b64_de_foto(f)
        mime = 'image/jpeg'
        tit  = sanitizar(f.get('titulo', f'Evidência {i}')).upper()
        desc = sanitizar(f.get('comentarios', 'N/A'))
        sev  = f.get('severidade', 'Normal')
        prazo = f.get('prazo_correcao', 'Monitorar')
        
        cls_b = {'Crítico': 'badge-critico', 'Critico': 'badge-critico', 'Observação': 'badge-obs', 'Observacao': 'badge-obs'}.get(sev, 'badge-normal')
        cls_p = {'Imediato (0–24h)': 'badge-prazo-imediato', 'Urgente (até 7 dias)': 'badge-prazo-urgente', 'Planejado (até 30 dias)': 'badge-prazo-planejado'}.get(prazo, 'badge-prazo-monitorar')
        cat  = sanitizar(f.get('categoria', 'Geral'))
        
        sev_norm = sev.replace('ã','a').replace('Ã','A')
        badge_html = f'<span class="badge {cls_b}">{sanitizar(sev)}</span> <span class="badge {cls_p}">{sanitizar(prazo)}</span>' if sev_norm not in ('Normal', '') else f'<span class="badge {cls_p}">{sanitizar(prazo)}</span>'
        
        mats_list = f.get('materiais') or []
        if not mats_list and f.get('material_necessario','').strip():
            mats_list = [{'descricao': f.get('material_necessario',''), 'unidade':'un', 'quantidade':1, 'custo_unit':0.0}]
        
        mat_html_bloco = ''
        if mats_list and any(m.get('descricao','').strip() for m in mats_list):
            subtotal_ev = sum(m.get('quantidade',0)*m.get('custo_unit',0) for m in mats_list)
            rows_mat = "".join(
                "<tr><td style='padding:3px 8px;border:1px solid #e2e8f0;'>" + sanitizar(m.get('descricao','')) + "</td>"
                "<td style='padding:3px 8px;border:1px solid #e2e8f0;text-align:center;'>" + sanitizar(m.get('unidade','un')) + "</td>"
                "<td style='padding:3px 8px;border:1px solid #e2e8f0;text-align:center;'>" + str(int(m.get('quantidade',1))) + "</td>"
                "<td style='padding:3px 8px;border:1px solid #e2e8f0;text-align:right;'>R$ " + f"{m.get('custo_unit',0):.2f}" + "</td>"
                "<td style='padding:3px 8px;border:1px solid #e2e8f0;text-align:right;font-weight:bold;'>R$ " + f"{m.get('quantidade',0)*m.get('custo_unit',0):.2f}" + "</td></tr>"
                for m in mats_list if m.get('descricao','').strip()
            )
            subtotal_html = (
                "<tr style='background:#fef2f2;'><td colspan='4' style='padding:3px 8px;border:1px solid #e2e8f0;text-align:right;font-weight:bold;color:#DA291C;'>Subtotal</td>"
                "<td style='padding:3px 8px;border:1px solid #e2e8f0;text-align:right;font-weight:bold;color:#DA291C;'>R$ " + f"{subtotal_ev:.2f}" + "</td></tr>"
            ) if subtotal_ev > 0 else ""
            mat_html_bloco = (
                '<div class="label-desc" style="margin-top:8px;color:#DA291C;">🔧 Materiais para Correção:</div>'
                '<table style="width:100%;border-collapse:collapse;font-size:8pt;margin-top:4px;">'
                '<tr style="background:#fef2f2;"><th style="padding:3px 8px;border:1px solid #e2e8f0;text-align:left;">Material</th>'
                '<th style="padding:3px 8px;border:1px solid #e2e8f0;">Un.</th>'
                '<th style="padding:3px 8px;border:1px solid #e2e8f0;">Qtd.</th>'
                '<th style="padding:3px 8px;border:1px solid #e2e8f0;">Unit. R$</th>'
                '<th style="padding:3px 8px;border:1px solid #e2e8f0;">Total R$</th></tr>'
                + rows_mat + subtotal_html + '</table>'
            )

        fotos_html += f"""
        <table class="card-evidencia">
          <tr>
            <td class="card-num">{i:02d}</td>
            <td class="col-foto"><img src="data:{mime};base64,{b64}"/></td>
            <td class="col-texto">
              <div class="foto-titulo">{tit}</div>
              {badge_html}
              <span style="font-size:7.5pt;color:{COR_CINZA};margin-left:6px;">{cat}</span>
              <div class="label-desc" style="margin-top:6px;">Descrição Técnica:</div>
              <div class="foto-desc">{desc}</div>
              {mat_html_bloco}
            </td>
          </tr>
        </table>"""

    extras_html = ""
    if extras:
        extras_html = '<div style="page-break-before:always;"></div><div class="section-header">4 &nbsp; ANEXOS ADICIONAIS</div>'
        for i, f in enumerate(extras, 1):
            b64  = _obter_b64_de_foto(f)
            mime = sanitizar(f.get('type', 'image/jpeg'))
            tit  = sanitizar(f.get('titulo', f'Anexo {i}')).upper()
            desc = sanitizar(f.get('comentarios', 'N/A'))
            extras_html += f"""
            <table class="card-evidencia">
              <tr>
                <td class="card-num">A{i:02d}</td>
                <td class="col-foto"><img src="data:{mime};base64,{b64}"/></td>
                <td class="col-texto">
                  <div class="foto-titulo">{tit}</div>
                  <div class="label-desc">Descrição / Contexto:</div>
                  <div class="foto-desc">{desc}</div>
                </td>
              </tr>
            </table>"""

    status = sanitizar(dados.get('status_laudo', 'N/I'))
    cor_status = {"✅ Aprovado": "#16A34A", "⚠️ Aprovado com Ressalvas": "#D97706", "❌ Reprovado": "#DA291C", "🔄 Em Acompanhamento": "#002060"}.get(status, "#1E293B")
    
    conclusao_html = f"""
    <div class="section-header">5 &nbsp; CONCLUSÃO E PARECER TÉCNICO</div>
    <table class="dados-table">
      <tr><td class="label">STATUS DO LAUDO</td><td class="value" colspan="3" style="font-weight:bold;color:{cor_status};">{status}</td></tr>
    </table>
    <div class="foto-desc" style="margin-bottom:28px;font-size:10pt;background:#F8FAFC;">
      {sanitizar(dados.get('conclusao', 'Sem parecer emitido.'))}
    </div>
    """

    qr_b64 = _gerar_qrcode_b64(f"{dados.get('numero_relatorio', 'GIRCP')} - {dados.get('site_id', '')}")
    qr_html = f'<img src="data:image/png;base64,{qr_b64}" style="width:65px; margin-bottom:-15px;"/>'

    html_raw = f"""<!DOCTYPE html>
<html lang="pt-BR"><head><meta charset="utf-8"><style>{_css_pdf()}</style></head>
<body>
{wm_html}
<div class="page-header">
  <div><div class="header-acento"></div><div class="header-titulo">Relatório de Visita Técnica</div></div>
  <div class="header-data">Data da Vistoria:<br><strong>{sanitizar(dados.get('data_hora',''))}</strong></div>
</div>
<div class="section-header">1 &nbsp; DADOS CADASTRAIS DA INFRAESTRUTURA</div>
<table class="dados-table">
  <tr><td class="label">N.º DO RELATÓRIO</td><td class="value">{sanitizar(dados.get('numero_relatorio','-'))}</td><td class="label">REVISÃO</td><td class="value">{sanitizar(dados.get('revisao','Rev.00'))}</td></tr>
  <tr><td class="label">SITE / IDENTIFICAÇÃO</td><td class="value">{sanitizar(dados.get('site_id',''))}</td><td class="label">ART / RRT Nº</td><td class="value">{sanitizar(dados.get('art_rrt','-'))}</td></tr>
  <tr><td class="label">TÍTULO</td><td class="value">{sanitizar(dados.get('titulo',''))}</td><td class="label">EMPRESA</td><td class="value">{sanitizar(dados.get('empresa',''))}</td></tr>
  <tr><td class="label">ENDEREÇO FÍSICO</td><td class="value" colspan="3">{sanitizar(dados.get('endereco',''))}</td></tr>
  <tr><td class="label">TÉCNICO EM CAMPO</td><td class="value" colspan="3">{sanitizar(dados.get('tecnico', dados.get('contato','')))}</td></tr>
  <tr><td class="label">CONTATO TÉCNICO</td><td class="value">{sanitizar(dados.get('contato',''))}</td><td class="label">TELEFONE</td><td class="value">{sanitizar(dados.get('telefone',''))}</td></tr>
</table>
{resumo_html}
<div class="section-header">3 &nbsp; REGISTRO FOTOGRÁFICO E EVIDÊNCIAS</div>
{fotos_html}
{extras_html}
{conclusao_html}
<div class="assinatura-wrapper">{qr_html}<br>{sig_img}<div class="assinatura-linha"></div><div class="assinatura-nome">{sanitizar(dados.get('contato','Responsável Técnico'))}</div><div class="assinatura-cargo">Responsável Técnico</div><div class="logo-wrapper">{logo_img}</div></div>
</body></html>"""

    nome = f"Relatorio_{sanitizar(dados.get('site_id','SITE')).replace(' ','_')}_{datetime.now().strftime('%Y%m%d%H%M%S')}.pdf"
    
    pdf_bytes_temp = HTML(string=html_raw.replace("HASH_PLACEHOLDER", "Gerando...")).write_pdf()
    hash_doc = hashlib.sha256(pdf_bytes_temp).hexdigest()
    
    html_final = html_raw.replace("HASH_PLACEHOLDER", f"SHA-256: {hash_doc[:16]}")
    pdf_bytes_final = HTML(string=html_final).write_pdf()
    
    return pdf_bytes_final, nome

def aplicar_estilo():
    st.markdown(f"""<style>
    [data-testid="stAppViewContainer"] {{ background: #F8FAFC; }}
    [data-testid="stSidebar"] {{ background: {COR_AZUL} !important; }}
    [data-testid="stSidebar"] * {{ color: #fff !important; }}
    .eng-banner {{ background: linear-gradient(135deg, {COR_AZUL} 0%, {COR_AZUL_MED} 100%); color: #fff; padding: 20px 28px; border-radius: 10px; margin-bottom: 24px; display: flex; align-items: center; justify-content: space-between; box-shadow: 0 4px 16px rgba(0,32,96,0.18); }}
    .eng-banner-title {{ font-size: 22px; font-weight: 900; letter-spacing: 1.2px; text-transform: uppercase; }}
    .eng-banner-badge {{ background: {COR_VERMELHO}; color: #fff; padding: 5px 14px; border-radius: 20px; font-size: 11px; font-weight: 700; letter-spacing: 0.5px; }}
    .eng-section {{ background: {COR_AZUL}; color: #fff; padding: 10px 18px; border-radius: 6px; margin: 24px 0 14px 0; font-weight: 700; font-size: 13px; text-transform: uppercase; border-left: 5px solid {COR_VERMELHO}; }}
    .eng-metric {{ background: #fff; border: 1.5px solid {COR_BORDA}; border-radius: 8px; padding: 14px 10px; text-align: center; }}
    .eng-metric-val {{ font-size: 26px; font-weight: 900; color: {COR_AZUL}; line-height: 1.1; }}
    .eng-metric-label {{ font-size: 10px; color: {COR_CINZA}; text-transform: uppercase; margin-top: 4px; letter-spacing: 0.5px; }}
    </style>""", unsafe_allow_html=True)

def banner(subtitulo: str = ""):
    subtitulo_seguro = sanitizar(subtitulo)
    sub_html = f" &mdash; {subtitulo_seguro}" if subtitulo_seguro else ""
    st.markdown(f"""
    <div class="eng-banner">
      <div>
        <div class="eng-banner-title">⚡ GIRCP</div>
        <div style="font-size: 12px; opacity: 0.75; margin-top: 2px;">GERADOR INTELIGENTE DE RELATÓRIOS E CONTROLE FOTOGRÁFICO{sub_html}</div>
      </div>
      <div class="eng-banner-badge">SISTEMA CORPORATIVO</div>
    </div>""", unsafe_allow_html=True)
    _botao_backup_db()

def secao(icone: str, titulo: str):
    st.markdown(f'<div class="eng-section">{sanitizar(icone)} &nbsp; {sanitizar(titulo)}</div>', unsafe_allow_html=True)

def _parse_kml_to_dataframe(arquivo_kml):
    if hasattr(arquivo_kml, 'seek'):
        arquivo_kml.seek(0)
        
    tree = ET.parse(arquivo_kml)
    root = tree.getroot()
    
    for elem in root.iter():
        if '}' in elem.tag:
            elem.tag = elem.tag.split('}', 1)[1]
            
    dados = []
    for placemark in root.findall('.//Placemark'):
        name_node = placemark.find('.//name')
        site_id = name_node.text.strip() if name_node is not None and name_node.text else "SEM NOME"
        
        lat, lon = 0.0, 0.0
        
        coords_node = placemark.find('.//coordinates')
        if coords_node is not None and coords_node.text:
            coords_str = coords_node.text.strip().split()
            if coords_str:
                valores = coords_str[0].split(',')
                try:
                    lon, lat = float(valores[0]), float(valores[1])
                except (ValueError, IndexError):
                    pass

        ext_data = placemark.find('.//ExtendedData')
        endereco, grupo = "", ""
        if ext_data is not None:
            for data in ext_data.findall('.//Data'):
                n_attr = str(data.get('name', '')).upper()
                v_node = data.find('value')
                val = v_node.text.strip() if v_node is not None and v_node.text else ""
                
                if 'ENDERE' in n_attr or 'RUA' in n_attr: endereco = val
                elif n_attr in ['GRUPO', 'ÁREA', 'AREA']: grupo = val
                elif 'LATITUDE' in n_attr or 'LAT' == n_attr:
                    try: lat = float(val.replace(',', '.'))
                    except ValueError: pass
                elif 'LONGITUDE' in n_attr or 'LON' == n_attr:
                    try: lon = float(val.replace(',', '.'))
                    except ValueError: pass
                    
        dados.append({'SITE': site_id, 'ENDEREÇO': endereco, 'GRUPO': grupo, 'LATITUDE': lat, 'LONGITUDE': lon})
    return pd.DataFrame(dados)

def _carregar_base_dados(arquivo):
    if not arquivo: return None
    try:
        if arquivo.name.lower().endswith('.kml'): return _parse_kml_to_dataframe(arquivo)
        else: return pd.read_excel(arquivo)
    except Exception as e:
        st.error(f"Erro ao processar a base: {e}")
        return None

def _obter_filtros_cascata(df_sites):
    if df_sites is None or df_sites.empty: return "", "", 0.0, 0.0
    col_site = next((col for col in df_sites.columns if str(col).upper() in ['SITE', 'SITES', 'IDENTIFICAÇÃO', 'NOME DO SITE']), None)
    if not col_site: return "", "", 0.0, 0.0
    col_grupo = next((col for col in df_sites.columns if 'GRUPO' in str(col).upper() or 'ÁREA' in str(col).upper() or 'AREA' in str(col).upper()), None)
    
    df_filtrado = df_sites
    c1, c2 = st.columns(2)
    with c1:
        if col_grupo:
            grupos = ["-- Todos --"] + sorted(df_sites[col_grupo].dropna().astype(str).unique().tolist())
            grupo_sel = st.selectbox("📌 1. Filtrar por Grupo/Tecnologia:", grupos, key="filtro_grupo_spc")
            if grupo_sel != "-- Todos --": df_filtrado = df_sites[df_sites[col_grupo] == grupo_sel]
    with c2:
        lista_sites = df_filtrado[col_site].dropna().astype(str).unique().tolist()
        escolha = st.selectbox("📌 2. Selecione o Site:", [OPT_DIGITAR_MANUAL] + sorted(lista_sites), key="filtro_escolha_site")
    
    if escolha == OPT_DIGITAR_MANUAL: return "", "", 0.0, 0.0
    col_end = next((col for col in df_sites.columns if 'ENDERE' in str(col).upper() or 'RUA' in str(col).upper()), None)
    endereco_val, lat_val, lon_val = "", 0.0, 0.0
    
    linha = df_sites[df_sites[col_site] == escolha]
    if not linha.empty: 
        if col_end: endereco_val = str(linha.iloc[0][col_end])
        if 'LATITUDE' in linha.columns: lat_val = float(linha.iloc[0]['LATITUDE'])
        if 'LONGITUDE' in linha.columns: lon_val = float(linha.iloc[0]['LONGITUDE'])
            
    return escolha, endereco_val, lat_val, lon_val

def _salvar_novo_relatorio(dados_cad, fotos, extras):
    with sqlite3.connect(DB_NAME) as conn:
        num_rel = proximo_numero_relatorio(conn)
        conn.execute(
            '''INSERT INTO relatorios
               (titulo, contato, empresa, telefone, email, site_id, endereco,
                data_hora, fotos_json, extras_json, tecnico, numero_relatorio, revisao, latitude, longitude, art_rrt, conclusao, status_laudo)
               VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)''',
            (sanitizar(dados_cad['titulo']), sanitizar(dados_cad['contato']),
             sanitizar(dados_cad['empresa']), sanitizar(dados_cad['telefone']),
             sanitizar(dados_cad['email']), sanitizar(dados_cad['site_id']),
             sanitizar(dados_cad['endereco']), sanitizar(dados_cad['data_hora']),
             json.dumps(fotos), json.dumps(extras),
             sanitizar(dados_cad['tecnico']), num_rel, 'Rev.00', dados_cad['latitude'], dados_cad['longitude'],
             sanitizar(dados_cad['art_rrt']), sanitizar(dados_cad['conclusao']), sanitizar(dados_cad['status_laudo'])))
    st.success(f"✅ RELATÓRIO **{sanitizar(dados_cad['site_id'])}** SALVO! ACESSE A ABA PARA GERAR PDF.")
    st.balloons()

def tela_novo():
    banner("NOVO RELATÓRIO")
    if "ordem_evidencias" not in st.session_state: st.session_state["ordem_evidencias"] = []

    st.info("💡 **Inteligência Analítica:** Importe arquivo **.KML** ou **.XLSX** para criar filtro em cascata.")
    arquivo_base = st.file_uploader("Importar Base de Sites (Opcional)", type=["kml", "xlsx", "xls"])
    df_sites = _carregar_base_dados(arquivo_base)
    site_selecionado, endereco_autofill, lat_autofill, lon_autofill = _obter_filtros_cascata(df_sites)

    if "site_anterior" not in st.session_state:
        st.session_state["site_anterior"] = None

    if st.session_state["site_anterior"] != site_selecionado:
        st.session_state["novo_lat"] = float(lat_autofill)
        st.session_state["novo_lon"] = float(lon_autofill)
        st.session_state["site_anterior"] = site_selecionado

    secao("📸", "2. EVIDÊNCIAS FOTOGRÁFICAS PRINCIPAIS")
    arq_fotos = st.file_uploader("FOTOS QUE DOCUMENTAM INTERVENÇÕES", type=["jpg", "jpeg", "png", "webp", "gif"], accept_multiple_files=True, key="up_evidencias_principal")
    
    if arq_fotos:
        atuais = [f.name for f in arq_fotos]
        st.session_state["ordem_evidencias"] = [f for f in st.session_state["ordem_evidencias"] if f in atuais]
        for f in atuais:
            if f not in st.session_state["ordem_evidencias"]:
                st.session_state["ordem_evidencias"].append(f)
    else: st.session_state["ordem_evidencias"] = []

    if arq_fotos and len(st.session_state["ordem_evidencias"]) > 1:
        st.caption("🔀 Reordene as fotos antes de preencher os metadados:")
        for idx, filename in enumerate(st.session_state["ordem_evidencias"]):
            c_n, c_u, c_d = st.columns([10, 1, 1])
            c_n.write(f"`{idx+1}.` {sanitizar(filename)}")
            total = len(st.session_state["ordem_evidencias"])
            if idx > 0 and c_u.button("⬆️", key=f"up_ev_{idx}"):
                o = st.session_state["ordem_evidencias"]
                o[idx], o[idx-1] = o[idx-1], o[idx]; st.rerun()
            if idx < total - 1 and c_d.button("⬇️", key=f"dw_ev_{idx}"):
                o = st.session_state["ordem_evidencias"]
                o[idx], o[idx+1] = o[idx+1], o[idx]; st.rerun()
        st.divider()

    secao("📋", "1. IDENTIFICAÇÃO DA INFRAESTRUTURA")
    c1, c2, c3 = st.columns(3)
    with c1:
        titulo  = st.text_input(LBL_TITULO, value="", placeholder="Relatório Técnico", key="novo_titulo")
        contato = st.text_input("CONTATO / TÉCNICO", value=st.session_state.get("_tecnico_global", ""), key="novo_contato")
    with c2:
        empresa  = st.text_input("EMPRESA", value="", placeholder="Empresa Parceira", key="novo_empresa")
        telefone = st.text_input("TELEFONE", value="", placeholder="(11) 99999-9999", key="novo_telefone")
    with c3:
        email   = st.text_input("E-MAIL", value="", placeholder="tecnico@empresa.com.br", key="novo_email")
        site_id = st.text_input("IDENTIFICAÇÃO DO SITE", value=site_selecionado)

    tecnico  = st.text_input("TÉCNICO EM CAMPO", value=st.session_state.get("_tecnico_global", ""), key="novo_tecnico")
    art_rrt = st.text_input("ART / RRT Nº", value="", placeholder="Ex: 2026/000123 — CREA-SP", key="novo_art_rrt")
    endereco = st.text_input("ENDEREÇO FÍSICO", value=endereco_autofill)

    c_lat_novo, c_lon_novo = st.columns(2)
    with c_lat_novo:
        latitude_manual = st.number_input("LATITUDE", format="%.6f", step=0.000001, key="novo_lat")
    with c_lon_novo:
        longitude_manual = st.number_input("LONGITUDE", format="%.6f", step=0.000001, key="novo_lon")

    col_d, col_h = st.columns(2)
    with col_d: data_vis = st.date_input("DATA DA VISITA", value=datetime.today(), key="novo_data")
    with col_h: hora_vis = st.time_input("HORA", value=datetime.now().time(), key="novo_hora")
    data_hora = f"{data_vis.strftime('%d/%m/%Y')} às {hora_vis.strftime('%H:%M')}"

    st.divider()
    secao("📸", "2. METADADOS DAS EVIDÊNCIAS")

    fotos_proc = []
    if arq_fotos and st.session_state["ordem_evidencias"]:
        dict_arquivos = {a.name: a for a in arq_fotos}
        for idx, filename in enumerate(st.session_state["ordem_evidencias"]):
            arquivo  = dict_arquivos[filename]
            raw      = arquivo.getvalue()
            
            tipo_real = filetype.guess(raw)
            if not tipo_real or tipo_real.extension not in ('jpg', 'jpeg', 'png', 'gif', 'webp'):
                st.error(f"Arquivo '{arquivo.name}' não é imagem. Upload rejeitado.")
                continue

            raw_comp = comprimir_para_pdf(raw)
            foto_id  = hashlib.sha256(raw).hexdigest()[:16]
            safe_key = f"ev_{idx}_{hashlib.md5(filename.encode()).hexdigest()[:8]}"
            
            caminho_foto = os.path.join(FOTOS_DIR, f"ev_{foto_id}.jpg")
            with open(caminho_foto, "wb") as f_out: f_out.write(raw_comp)

            st.markdown("---")
            c_img, c_dados = st.columns([1, 3])
            with c_img:
                st.image(arquivo, use_container_width=True)
                st.caption(f"ID: {sanitizar(foto_id)}")
            with c_dados:
                tit = st.text_input(LBL_TITULO, key=f"t_{safe_key}")
                com = st.text_area(LBL_DESCRICAO, key=f"c_{safe_key}", height=75)
                c_sev, c_cat, c_prazo = st.columns(3)
                with c_sev: sev = st.selectbox("SEVERIDADE", ["Normal", "Observacao", "Critico"], key=f"sev_{safe_key}")
                with c_cat: cat = st.selectbox("CATEGORIA",  ["Geral", "Antes", "Depois", "Detalhe"], key=f"cat_{safe_key}")
                with c_prazo: prazo = st.selectbox("PRAZO", ["Imediato (0–24h)", "Urgente (até 7 dias)", "Planejado (até 30 dias)", "Monitorar"], key=f"prazo_{safe_key}")
                
                st.markdown("**🔧 Materiais necessários**")
                mat_list_key = f"mats_{safe_key}"
                if mat_list_key not in st.session_state:
                    st.session_state[mat_list_key] = [{"descricao": "", "unidade": "un", "quantidade": 1, "custo_unit": 0.0}]
                for mi, mitem in enumerate(st.session_state[mat_list_key]):
                    mc1, mc2, mc3, mc4, mc5 = st.columns([3, 1.2, 1, 1.4, 0.5])
                    mitem["descricao"]  = mc1.text_input("Descrição", value=mitem["descricao"],  key=f"md_{safe_key}_{mi}", label_visibility="collapsed", placeholder="Ex: Disjuntor 40A")
                    mitem["unidade"]    = mc2.selectbox("Un.", ["un","m","kg","kit","cx","hr"], index=["un","m","kg","kit","cx","hr"].index(mitem["unidade"]), key=f"mu_{safe_key}_{mi}", label_visibility="collapsed")
                    mitem["quantidade"] = mc3.number_input("Qtd", value=float(mitem["quantidade"]), min_value=0.0, step=1.0, key=f"mq_{safe_key}_{mi}", label_visibility="collapsed")
                    mitem["custo_unit"] = mc4.number_input("R$ unit.", value=float(mitem["custo_unit"]), min_value=0.0, step=0.01, format="%.2f", key=f"mc_{safe_key}_{mi}", label_visibility="collapsed")
                    if mc5.button("❌", key=f"mdel_{safe_key}_{mi}") and len(st.session_state[mat_list_key]) > 1:
                        st.session_state[mat_list_key].pop(mi); st.rerun()
                if st.button("＋ Adicionar material", key=f"madd_{safe_key}"):
                    st.session_state[mat_list_key].append({"descricao": "", "unidade": "un", "quantidade": 1, "custo_unit": 0.0}); st.rerun()
                subtotal = sum(m["quantidade"] * m["custo_unit"] for m in st.session_state[mat_list_key])
                if subtotal > 0:
                    st.markdown(f"<div style='text-align:right;font-size:12px;color:{COR_AZUL};'>Subtotal: <strong>R$ {subtotal:,.2f}</strong></div>", unsafe_allow_html=True)

                fotos_proc.append({
                    "foto_id": foto_id, "caminho": caminho_foto, "type": "image/jpeg",
                    "titulo": tit.strip() or f"Evidência {idx+1}", "comentarios": com.strip() or "N/A",
                    "filename": arquivo.name, "severidade": sev, "categoria": cat, "prazo_correcao": prazo,
                    "materiais": st.session_state.get(mat_list_key, []),
                    "material_necessario": ", ".join(m["descricao"] for m in st.session_state.get(mat_list_key, []) if m["descricao"].strip()),
                })

    secao("📎", "3. ANEXOS ADICIONAIS")
    arq_extras = st.file_uploader("SITUAÇÃO ANTERIOR OU CONTEXTO GERAL", type=["jpg", "jpeg", "png", "webp", "gif"], accept_multiple_files=True, key="up_extras_principal")
    extras_proc = []
    if arq_extras:
        for idx_ex, arq_ex in enumerate(arq_extras):
            raw_ex = arq_ex.getvalue()
            tipo_real = filetype.guess(raw_ex)
            if not tipo_real or tipo_real.extension not in ('jpg', 'jpeg', 'png', 'gif', 'webp'):
                st.error(f"Arquivo '{arq_ex.name}' não é imagem. Upload rejeitado.")
                continue

            raw_comp_ex = comprimir_para_pdf(raw_ex)
            fid_ex = hashlib.sha256(raw_ex).hexdigest()[:16]
            
            caminho_extra = os.path.join(FOTOS_DIR, f"ex_{fid_ex}.jpg")
            with open(caminho_extra, "wb") as f_out_ex: f_out_ex.write(raw_comp_ex)
            
            c_img_ex, c_dados_ex = st.columns([1, 3])
            with c_img_ex: st.image(arq_ex, use_container_width=True)
            with c_dados_ex:
                t_ex = st.text_input(LBL_TITULO, key=f"t_ex_{idx_ex}")
                c_ex = st.text_area(LBL_DESCRICAO, key=f"c_ex_{idx_ex}", height=75)
                extras_proc.append({
                    "foto_id": fid_ex, "caminho": caminho_extra, "type": "image/jpeg",
                    "titulo": t_ex.strip() or f"Anexo {idx_ex+1}", "comentarios": c_ex.strip() or "N/A",
                    "filename": arq_ex.name
                })
    
    secao("📝", "4. CONCLUSÃO E PARECER TÉCNICO")
    status_laudo = st.selectbox("STATUS GERAL DO LAUDO", ["✅ Aprovado", "⚠️ Aprovado com Ressalvas", "❌ Reprovado", "🔄 Em Acompanhamento"], key="novo_status_laudo")
    conclusao = st.text_area("PARECER TÉCNICO / CONCLUSÃO", height=120, placeholder="Descreva o resultado geral da vistoria, recomendações...", key="novo_conclusao")

    dados_cad = {
        "titulo": titulo, "contato": contato, "empresa": empresa, "telefone": telefone,
        "email": email, "site_id": site_id, "endereco": endereco, "data_hora": data_hora,
        "tecnico": tecnico or contato, "latitude": latitude_manual, "longitude": longitude_manual,
        "art_rrt": art_rrt, "conclusao": conclusao, "status_laudo": status_laudo
    }

    c_prev, c_sub = st.columns([1, 2])
    with c_prev:
        if st.button("👁️ PRÉ-VISUALIZAR PDF", use_container_width=True):
            pdf_bytes_tmp, _ = gerar_pdf(dados_cad, fotos_proc, extras_proc)
            st.download_button("⬇️ Baixar Preview", pdf_bytes_tmp, "Preview.pdf", "application/pdf")
            
    with c_sub:
        submit = st.button("💾 SALVAR RELATÓRIO OFICIAL", type="primary", use_container_width=True)

    if submit: 
        erros = []
        if not dados_cad["site_id"].strip(): erros.append("• A Identificação do Site é obrigatória.")
        if not dados_cad["tecnico"].strip() and not dados_cad["contato"].strip(): erros.append("• Informe o Técnico em Campo.")
        if not fotos_proc: erros.append("• Adicione ao menos uma evidência.")
        
        if erros:
            st.error("⚠️ Corrija para salvar:\n" + "\n".join(erros))
        else:
            _salvar_novo_relatorio(dados_cad, fotos_proc, extras_proc)
            registrar_auditoria("criar", detalhe=dados_cad["site_id"])

def _render_cadastrais(row, lid):
    secao("📋", "DADOS CADASTRAIS")
    e1, e2, e3 = st.columns(3)
    with e1:
        tit = st.text_input(LBL_TITULO, value=row['titulo'], key=f"tit_{lid}")
        con = st.text_input("CONTATO", value=row['contato'], key=f"con_{lid}")
        
        val_art = row['art_rrt'] if 'art_rrt' in row.keys() and row['art_rrt'] else ""
        art = st.text_input("ART / RRT Nº", value=val_art, key=f"art_{lid}")
    with e2:
        emp = st.text_input("EMPRESA", value=row['empresa'], key=f"emp_{lid}")
        tel = st.text_input("TELEFONE", value=row['telefone'], key=f"tel_{lid}")
        
        val_status = row['status_laudo'] if 'status_laudo' in row.keys() and row['status_laudo'] else "✅ Aprovado"
        status = st.selectbox("STATUS DO LAUDO", ["✅ Aprovado", "⚠️ Aprovado com Ressalvas", "❌ Reprovado", "🔄 Em Acompanhamento"], index=["✅ Aprovado", "⚠️ Aprovado com Ressalvas", "❌ Reprovado", "🔄 Em Acompanhamento"].index(val_status), key=f"status_{lid}")
    with e3:
        eml = st.text_input("E-MAIL", value=row['email'], key=f"eml_{lid}")
        sit = st.text_input("SITE", value=row['site_id'], key=f"sit_{lid}")
        dat = st.text_input("DATA E HORA", value=row['data_hora'], key=f"dat_{lid}")
    
    val_end = row['endereco'] if 'endereco' in row.keys() and row['endereco'] else ""
    end = st.text_input("ENDEREÇO", value=val_end, key=f"end_{lid}")
    
    val_conc = row['conclusao'] if 'conclusao' in row.keys() and row['conclusao'] else ""
    conclusao = st.text_area("CONCLUSÃO E PARECER", value=val_conc, key=f"conclusao_{lid}")
    
    c_lat, c_lon = st.columns(2)
    with c_lat:
        lat_val = float(row['latitude']) if 'latitude' in row.keys() and row['latitude'] else 0.0
        lat = st.number_input("LATITUDE", value=lat_val, format="%.6f", step=0.000001, key=f"lat_{lid}")
    with c_lon:
        lon_val = float(row['longitude']) if 'longitude' in row.keys() and row['longitude'] else 0.0
        lon = st.number_input("LONGITUDE", value=lon_val, format="%.6f", step=0.000001, key=f"lon_{lid}")

    return {"tit": tit, "con": con, "emp": emp, "tel": tel, "eml": eml, "sit": sit, "dat": dat, "end": end, "lat": lat, "lon": lon, "art_rrt": art, "status_laudo": status, "conclusao": conclusao}

def _executar_acao_inline(lid, fid, acao, prefixo, db_field):
    if db_field not in ("fotos_json", "extras_json"):
        return
    
    _QUERIES_CAMPO = {
        "fotos_json": {
            "select": "SELECT fotos_json FROM relatorios WHERE id=?",
            "update": "UPDATE relatorios SET fotos_json=? WHERE id=?"
        },
        "extras_json": {
            "select": "SELECT extras_json FROM relatorios WHERE id=?",
            "update": "UPDATE relatorios SET extras_json=? WHERE id=?"
        }
    }

    with sqlite3.connect(DB_NAME) as conn:
        row = conn.execute(_QUERIES_CAMPO[db_field]["select"], (lid,)).fetchone()
        if not row or not row[0]: return
        fotos = json.loads(row[0])
        for i in range(len(fotos)):
            cfid = fotos[i].get("foto_id", fotos[i].get("base64", "")[:16])
            t_val = st.session_state.get(f"{prefixo}t_{lid}_{cfid}")
            c_val = st.session_state.get(f"{prefixo}c_{lid}_{cfid}")
            m_val = st.session_state.get(f"{prefixo}m_{lid}_{cfid}")
            p_val = st.session_state.get(f"{prefixo}p_{lid}_{cfid}")
            if t_val is not None: fotos[i]["titulo"] = t_val
            if c_val is not None: fotos[i]["comentarios"] = c_val
            if m_val is not None: fotos[i]["material_necessario"] = m_val
            if p_val is not None: fotos[i]["prazo_correcao"] = p_val
        
        k = next((i for i, f in enumerate(fotos) if f.get("foto_id", f.get("base64", "")[:16]) == fid), -1)
        if k != -1:
            if acao == "up" and k > 0: fotos[k], fotos[k-1] = fotos[k-1], fotos[k]
            elif acao == "down" and k < len(fotos)-1: fotos[k], fotos[k+1] = fotos[k+1], fotos[k]
            elif acao == "del":
                caminho_del = fotos[k].get("caminho")
                if caminho_del and os.path.exists(caminho_del):
                    try: os.remove(caminho_del)
                    except: pass
                fotos.pop(k)
            conn.execute(_QUERIES_CAMPO[db_field]["update"], (json.dumps(fotos), lid))
            conn.commit()
            registrar_auditoria(f"editou_inline_{acao}", lid, fid)

def _render_item_edicao(f, k, lid, prefixo, total_fotos, db_field):
    fid = f.get("foto_id", f.get("base64", "")[:16])
    st.markdown(
        f'<div style="background:#fff;border:1px solid {COR_BORDA};border-left:4px solid {COR_AZUL};'
        f'border-radius:8px;padding:12px;margin-bottom:10px;">'
        f'<span style="background:{COR_AZUL};color:#fff;font-size:10px;font-weight:700;'
        f'padding:2px 8px;border-radius:12px;">FOTO {k+1:02d}</span>'
        f'&nbsp;<span style="font-size:9px;color:{COR_CINZA};font-family:monospace;">ID: {sanitizar(fid)}</span></div>',
        unsafe_allow_html=True
    )
    col_i, col_d, col_ctrl = st.columns([1, 3, 0.4])
    b64_data = _obter_b64_de_foto(f)
    uri = f"data:{sanitizar(f.get('type', 'image/jpeg'))};base64,{b64_data}"
    
    if b64_data: col_i.markdown(f'<img src="{uri}" style="width:100%;border-radius:6px;"/>', unsafe_allow_html=True)
    with col_d:
        nt = st.text_input(LBL_TITULO, value=f.get('titulo', ''), key=f"{prefixo}t_{lid}_{fid}")
        nc = st.text_area(LBL_DESCRICAO, value=f.get('comentarios', ''), key=f"{prefixo}c_{lid}_{fid}", height=65)
        nprazo = st.selectbox("PRAZO", ["Imediato (0–24h)", "Urgente (até 7 dias)", "Planejado (até 30 dias)", "Monitorar"], index=["Imediato (0–24h)", "Urgente (até 7 dias)", "Planejado (até 30 dias)", "Monitorar"].index(f.get("prazo_correcao", "Monitorar")), key=f"{prefixo}p_{lid}_{fid}")
        
        st.markdown("**🔧 Materiais necessários**")
        emat_key = f"emats_{prefixo}_{lid}_{fid}"
        mats_default = f.get('materiais') or ([{"descricao": f.get('material_necessario',''), "unidade":"un","quantidade":1,"custo_unit":0.0}] if f.get('material_necessario','').strip() else [{"descricao":"","unidade":"un","quantidade":1,"custo_unit":0.0}])
        if emat_key not in st.session_state:
            st.session_state[emat_key] = mats_default
        for mi, mitem in enumerate(st.session_state[emat_key]):
            ec1, ec2, ec3, ec4, ec5 = st.columns([3, 1.2, 1, 1.4, 0.5])
            mitem["descricao"]  = ec1.text_input("Descrição", value=mitem.get("descricao",""),  key=f"emd_{prefixo}_{lid}_{fid}_{mi}", label_visibility="collapsed", placeholder="Ex: Disjuntor 40A")
            mitem["unidade"]    = ec2.selectbox("Un.", ["un","m","kg","kit","cx","hr"], index=["un","m","kg","kit","cx","hr"].index(mitem.get("unidade","un")), key=f"emu_{prefixo}_{lid}_{fid}_{mi}", label_visibility="collapsed")
            mitem["quantidade"] = ec3.number_input("Qtd", value=float(mitem.get("quantidade",1)), min_value=0.0, step=1.0, key=f"emq_{prefixo}_{lid}_{fid}_{mi}", label_visibility="collapsed")
            mitem["custo_unit"] = ec4.number_input("R$ unit.", value=float(mitem.get("custo_unit",0.0)), min_value=0.0, step=0.01, format="%.2f", key=f"emc_{prefixo}_{lid}_{fid}_{mi}", label_visibility="collapsed")
            if ec5.button("❌", key=f"emdel_{prefixo}_{lid}_{fid}_{mi}") and len(st.session_state[emat_key]) > 1:
                st.session_state[emat_key].pop(mi); st.rerun()
        if st.button("＋ Adicionar material", key=f"emadd_{prefixo}_{lid}_{fid}"):
            st.session_state[emat_key].append({"descricao":"","unidade":"un","quantidade":1,"custo_unit":0.0}); st.rerun()
        esubtotal = sum(m.get("quantidade",0)*m.get("custo_unit",0) for m in st.session_state[emat_key])
        if esubtotal > 0:
            st.markdown(f"<div style='text-align:right;font-size:12px;color:{COR_AZUL};'>Subtotal: <strong>R$ {esubtotal:,.2f}</strong></div>", unsafe_allow_html=True)
        nm = ", ".join(m["descricao"] for m in st.session_state[emat_key] if m.get("descricao","").strip())
    with col_ctrl:
        st.markdown("<br>", unsafe_allow_html=True)
        if k > 0 and st.button("⬆️", key=f"{prefixo}up_{lid}_{fid}"):
            _executar_acao_inline(lid, fid, "up", prefixo, db_field); st.rerun()
        if k < total_fotos-1 and st.button("⬇️", key=f"{prefixo}dn_{lid}_{fid}"):
            _executar_acao_inline(lid, fid, "down", prefixo, db_field); st.rerun()
        if st.button("❌", key=f"{prefixo}del_{lid}_{fid}"):
            _executar_acao_inline(lid, fid, "del", prefixo, db_field); st.rerun()
    emat_key2 = f"emats_{prefixo}_{lid}_{fid}"
    fc = f.copy(); fc['titulo'] = nt; fc['comentarios'] = nc; fc['material_necessario'] = nm; fc['prazo_correcao'] = nprazo
    fc['materiais'] = st.session_state.get(emat_key2, f.get('materiais', []))
    return fc

def _render_edicao_lista(fotos, lid, titulo_sec, icone, prefixo):
    secao(icone, titulo_sec)
    editados = []
    db_field = "fotos_json" if prefixo == "f" else "extras_json"
    for k, f in enumerate(fotos):
        fc = _render_item_edicao(f, k, lid, prefixo, len(fotos), db_field)
        editados.append(fc)
        st.divider()
    return editados

def _render_novas_fotos(lid, db_existentes):
    secao("➕", "ADICIONAR NOVAS FOTOS")
    arq = st.file_uploader("ENVIAR NOVAS IMAGENS", type=['png', 'jpg', 'jpeg', 'webp', 'gif'], accept_multiple_files=True, key=f"new_{lid}")
    novas = []
    if arq:
        ids = {f.get("foto_id") for f in db_existentes}
        for idx_n, a in enumerate(arq):
            raw = a.getvalue()
            tipo_real = filetype.guess(raw)
            if not tipo_real or tipo_real.extension not in ('jpg', 'jpeg', 'png', 'gif', 'webp'):
                st.error(f"Arquivo '{a.name}' não é imagem. Upload rejeitado.")
                continue

            raw_comp = comprimir_para_pdf(raw)
            fid = hashlib.sha256(raw).hexdigest()[:16]
            caminho_foto = os.path.join(FOTOS_DIR, f"new_{lid}_{fid}.jpg")
            with open(caminho_foto, "wb") as f_out: f_out.write(raw_comp)
                
            col_i, col_d = st.columns([1, 4])
            with col_i: st.image(a, width=90); st.caption(f"ID: {sanitizar(fid)}")
            with col_d:
                if fid in ids: st.warning(f"⚠️ DUPLICATA IGNORADA: {sanitizar(a.name)}"); continue
                nt = st.text_input(LBL_TITULO, key=f"n_t_{lid}_{idx_n}")
                nc = st.text_area(LBL_DESCRICAO, key=f"n_c_{lid}_{idx_n}", height=55)
            novas.append({
                "foto_id": fid, "caminho": caminho_foto, "type": "image/jpeg", 
                "titulo": nt.strip() if nt else f"Nova Foto {idx_n+1}",
                "comentarios": nc.strip() if nc else "N/A", "filename": a.name
            })
            ids.add(fid)
    return novas

def _tratar_botoes_acao_pdf(lid, row, state):
    if st.button(f"📄 GERAR PDF — {sanitizar(row['site_id'])}", key=f"pdf_{lid}", type="primary"):
        num_rel = row['numero_relatorio'] if 'numero_relatorio' in row.keys() and row['numero_relatorio'] else ""
        if not num_rel.strip():
            with sqlite3.connect(DB_NAME) as conn:
                num_rel = proximo_numero_relatorio(conn)
                conn.execute("UPDATE relatorios SET numero_relatorio=?, revisao=? WHERE id=?", (num_rel, "Rev.00", lid))
        
        dados_pdf = {
            "titulo": row['titulo'], 
            "contato": row['contato'], 
            "empresa": row['empresa'], 
            "telefone": row['telefone'],
            "email": row['email'], 
            "site_id": row['site_id'], 
            "endereco": row['endereco'] if 'endereco' in row.keys() else '',
            "data_hora": row['data_hora'], 
            "tecnico": row['tecnico'] if 'tecnico' in row.keys() and row['tecnico'] else row['contato'],
            "numero_relatorio": num_rel, 
            "revisao": row['revisao'] if 'revisao' in row.keys() and row['revisao'] else 'Rev.00',
            "art_rrt": row['art_rrt'] if 'art_rrt' in row.keys() and row['art_rrt'] else '', 
            "conclusao": row['conclusao'] if 'conclusao' in row.keys() and row['conclusao'] else '', 
            "status_laudo": row['status_laudo'] if 'status_laudo' in row.keys() and row['status_laudo'] else '✅ Aprovado'
        }
        
        with st.spinner("GERANDO PDF EM MEMÓRIA..."):
            pdf_bytes, file_name = gerar_pdf(dados_pdf, state["fotos_db"], state["extras_db"])
        
        st.download_button("⬇️ BAIXAR PDF GERADO", data=pdf_bytes, file_name=file_name, mime="application/pdf", key=f"dl_{lid}")
        registrar_auditoria("gerar_pdf", lid)

def _tratar_botoes_acao(lid, state):
    col_b, col_c = st.columns(2)
    with col_b:
        if st.button("🗑️ LIMPAR TODAS AS FOTOS", key=f"lim_{lid}"): st.session_state[f"conf_lim_{lid}"] = True
    with col_c: _tratar_botoes_acao_pdf(lid, row=state["row"], state=state)

def _tratar_limpeza(lid):
    if st.session_state.get(f"conf_lim_{lid}"):
        st.warning("⚠️ TEM CERTEZA? ESTA AÇÃO APAGARÁ **TODAS** AS FOTOS E ANEXOS.")
        cs, cn = st.columns(2)
        if cs.button("✅ SIM, APAGAR TUDO", key=f"sim_{lid}"):
            with sqlite3.connect(DB_NAME) as conn:
                conn.execute("UPDATE relatorios SET fotos_json='[]', extras_json='[]' WHERE id=?", (lid,))
                conn.commit()
            registrar_auditoria("limpar_fotos", lid)
            st.session_state.pop(f"conf_lim_{lid}", None); st.rerun()
        if cn.button("❌ CANCELAR", key=f"nao_{lid}"):
            st.session_state.pop(f"conf_lim_{lid}", None); st.rerun()

def _salvar_edicoes(lid, state):
    fotos_finais = state["fotos_edit"] + state["novas"]
    extras_finais = state["extras_edit"] 
    d = state["cad"]
    
    if not d["sit"].strip():
        st.error("• A Identificação do Site é obrigatória.")
        return
        
    with sqlite3.connect(DB_NAME) as conn:
        conn.execute('''UPDATE relatorios
                        SET titulo=?,contato=?,empresa=?,telefone=?,email=?,
                            site_id=?,endereco=?,data_hora=?,
                            fotos_json=?,extras_json=?, latitude=?, longitude=?,
                            art_rrt=?, conclusao=?, status_laudo=? WHERE id=?''',
                     (d["tit"], d["con"], d["emp"], d["tel"], d["eml"],
                      d["sit"], d["end"], d["dat"],
                      json.dumps(fotos_finais), json.dumps(extras_finais), d["lat"], d["lon"],
                      d["art_rrt"], d["conclusao"], d["status_laudo"], lid))
        conn.commit()
    registrar_auditoria("editar", lid, d["sit"])
    st.success("✅ RELATÓRIO ATUALIZADO COM SUCESSO!"); st.rerun()

def _processar_acoes_relatorio(state: dict):
    lid = state["lid"]
    _tratar_botoes_acao(lid, state)
    _tratar_limpeza(lid)
    if state["salvar"]: _salvar_edicoes(lid, state)

def _render_dados_cadastrais_form(row, lid, fotos_db, extras_db):
    cad = _render_cadastrais(row, lid)
    fotos_edit = _render_edicao_lista(fotos_db, lid, "EDITAR EVIDÊNCIAS", "📸", "f")
    extras_edit = _render_edicao_lista(extras_db, lid, "EDITAR ANEXOS", "📎", "e")
    novas = _render_novas_fotos(lid, fotos_db + extras_db)
    salvar = st.button("🔄  SALVAR ALTERAÇÕES", type="primary", use_container_width=True, key=f"salvar_{lid}")
    return {"lid": lid, "row": row, "salvar": salvar, "fotos_db": fotos_db, "extras_db": extras_db, "fotos_edit": fotos_edit, "extras_edit": extras_edit, "novas": novas, "cad": cad}

def _render_relatorio_expander(row):
    lid = row['id']
    exp_key = f"exp_aberto_{lid}"
    if exp_key not in st.session_state:
        st.session_state[exp_key] = False

    aberto = st.session_state[exp_key]
    icone = "🔽" if aberto else "▶️"
    if st.button(f"{icone}  📍 {sanitizar(row['site_id'])}  |  {sanitizar(row['data_hora'])}  |  ID #{lid}", key=f"toggle_{lid}", use_container_width=True):
        st.session_state[exp_key] = not aberto
        st.rerun()

    if st.session_state[exp_key]:
        fotos_db = json.loads(row['fotos_json'] if row['fotos_json'] else "[]")
        
        ext_str = row['extras_json'] if 'extras_json' in row.keys() and row['extras_json'] else "[]"
        extras_db = json.loads(ext_str)
        
        with st.container(border=True):
            state = _render_dados_cadastrais_form(row, lid, fotos_db, extras_db)
            _processar_acoes_relatorio(state)

def tela_pesquisa():
    banner("PESQUISAR, EDITAR E EXPORTAR")
    secao("🔍", "BUSCAR RELATÓRIOS")
    col_busca, col_lim = st.columns([3, 1])
    with col_busca: termo = st.text_input("NOME DO SITE OU PARTE DO TÍTULO:", placeholder="Ex: SMSMT15")
    with col_lim: limite = st.selectbox("EXIBIR", [10, 25, 50, 100], index=0)

    with sqlite3.connect(DB_NAME) as conn:
        conn.row_factory = sqlite3.Row
        if termo.strip():
            rows = conn.execute("SELECT * FROM relatorios WHERE site_id LIKE ? OR titulo LIKE ? ORDER BY id DESC LIMIT ?", (f'%{termo}%', f'%{termo}%', limite)).fetchall()
        else:
            rows = conn.execute("SELECT * FROM relatorios ORDER BY id DESC LIMIT ?", (limite,)).fetchall()

    if not rows:
        st.info("NENHUM RELATÓRIO ENCONTRADO. CRIE UM NA ABA 'NOVO RELATÓRIO'.")
        return

    st.markdown("<br>", unsafe_allow_html=True)
    ta, tb, tc = st.columns(3)
    with ta: st.markdown(f'<div class="eng-metric"><div class="eng-metric-val">{len(rows)}</div><div class="eng-metric-label">RESULTADOS</div></div>', unsafe_allow_html=True)
    with tb: st.markdown(f'<div class="eng-metric"><div class="eng-metric-val">{sum(len(json.loads(r["fotos_json"] if r["fotos_json"] else "[]")) for r in rows)}</div><div class="eng-metric-label">EVIDÊNCIAS</div></div>', unsafe_allow_html=True)
    with tc: st.markdown(f'<div class="eng-metric"><div class="eng-metric-val">{sum(len(json.loads(r["extras_json"] if "extras_json" in r.keys() and r["extras_json"] else "[]")) for r in rows)}</div><div class="eng-metric-label">ANEXOS</div></div>', unsafe_allow_html=True)
    st.markdown("<br>", unsafe_allow_html=True)
    st.divider()

    for row in rows: _render_relatorio_expander(row)

def tela_dashboard():
    banner("DASHBOARD")
    with sqlite3.connect(DB_NAME) as conn:
        df = pd.read_sql_query("SELECT * FROM relatorios ORDER BY id DESC", conn)

    if df.empty:
        st.info("NENHUM RELATÓRIO CADASTRADO AINDA.")
        return

    df['tecnico'] = df['tecnico'].fillna(df['contato'])
    df['qtd_fotos'] = df['fotos_json'].apply(lambda x: len(json.loads(x or "[]")))
    df['qtd_extras'] = df.get('extras_json', pd.Series(['[]']*len(df))).apply(lambda x: len(json.loads(x or "[]")))
    df['total_imagens'] = df['qtd_fotos'] + df['qtd_extras']
    
    try:
        df['data_formatada'] = pd.to_datetime(df['data_hora'].str.extract(r'(\d{2}/\d{2}/\d{4})')[0], format='%d/%m/%Y', errors='coerce')
    except:
        df['data_formatada'] = pd.NaT

    st.markdown("### 🎛️ Filtros Analíticos")
    c_tec, c_site = st.columns(2)
    with c_tec:
        lista_tecnicos = df['tecnico'].dropna().unique().tolist()
        tec_sel = st.multiselect("Filtrar por Técnico:", lista_tecnicos, default=lista_tecnicos)
    with c_site:
        lista_sites = df['site_id'].dropna().unique().tolist()
        site_sel = st.multiselect("Filtrar por Site:", lista_sites, default=lista_sites)
    
    df_filtrado = df[(df['tecnico'].isin(tec_sel)) & (df['site_id'].isin(site_sel))]

    if df_filtrado.empty:
        st.warning("Nenhum dado corresponde aos filtros selecionados.")
        return

    secao("📊", "RESUMO GERAL")
    m1, m2, m3, m4 = st.columns(4)
    m1.markdown(f'<div class="eng-metric"><div class="eng-metric-val">{len(df_filtrado)}</div><div class="eng-metric-label">RELATÓRIOS</div></div>', unsafe_allow_html=True)
    m2.markdown(f'<div class="eng-metric"><div class="eng-metric-val">{df_filtrado["site_id"].nunique()}</div><div class="eng-metric-label">SITES ÚNICOS</div></div>', unsafe_allow_html=True)
    m3.markdown(f'<div class="eng-metric"><div class="eng-metric-val">{df_filtrado["qtd_fotos"].sum()}</div><div class="eng-metric-label">EVIDÊNCIAS</div></div>', unsafe_allow_html=True)
    m4.markdown(f'<div class="eng-metric"><div class="eng-metric-val">{df_filtrado["qtd_extras"].sum()}</div><div class="eng-metric-label">ANEXOS</div></div>', unsafe_allow_html=True)

    st.markdown("---")
    
    col_chart1, col_chart2 = st.columns(2)
    
    with col_chart1:
        secao("📈", "PRODUTIVIDADE (TOP SITES)")
        df_agrupado = df_filtrado.groupby("site_id")['total_imagens'].sum().reset_index().sort_values("total_imagens", ascending=False).head(10)
        if not df_agrupado.empty:
            fig_bar = px.bar(df_agrupado, x="site_id", y="total_imagens", text="total_imagens", color="total_imagens", 
                             color_continuous_scale=px.colors.sequential.Reds, labels={"site_id": "Site", "total_imagens": "Total de Imagens"})
            fig_bar.update_layout(margin=dict(l=0, r=0, t=30, b=0), showlegend=False)
            st.plotly_chart(fig_bar, use_container_width=True)

    with col_chart2:
        secao("🚨", "DISTRIBUIÇÃO DE SEVERIDADE")
        severidades = {"Critico": 0, "Observacao": 0, "Normal": 0}
        
        for _, r in df_filtrado.iterrows():
            for foto in json.loads(r['fotos_json'] or "[]"):
                sev = foto.get('severidade', 'Normal')
                sev_norm = {'Crítico': 'Critico', 'Observação': 'Observacao'}.get(sev, sev)
                severidades[sev_norm] = severidades.get(sev_norm, 0) + 1
        
        df_sev = pd.DataFrame(list(severidades.items()), columns=['Severidade', 'Quantidade'])
        if df_sev['Quantidade'].sum() > 0:
            fig_pie = px.pie(df_sev, values='Quantidade', names='Severidade', hole=0.4, 
                             color='Severidade', color_discrete_map={"Critico": COR_VERMELHO, "Observacao": COR_AMARELO, "Normal": COR_VERDE})
            fig_pie.update_layout(margin=dict(l=0, r=0, t=30, b=0))
            st.plotly_chart(fig_pie, use_container_width=True)
        else:
            st.info("Sem dados de severidade para exibir.")

    st.markdown("---")
    
    secao("🌍", "MAPA TÁTICO DE VISTORIAS (GEOLOCALIZAÇÃO)")
    col_st1, col_st2 = st.columns([2, 1])
    with col_st1:
        status_opcoes = ["Todos", "✅ Concluída", "⚙️ Em Andamento", "🕐 Pendente"]
        status_filtro = st.radio("Exibir visitas:", status_opcoes, horizontal=True, key="mapa_status_filtro")
    with col_st2:
        st.markdown(
            f"""<div style='background:#fff;border:1px solid {COR_BORDA};border-radius:8px;padding:10px 14px;font-size:11px;line-height:1.8;'>
            <span style='color:#16A34A;font-size:15px;'>●</span> <b>Concluída</b> — relatório com fotos<br>
            <span style='color:#D97706;font-size:15px;'>●</span> <b>Em Andamento</b> — sem fotos ainda<br>
            <span style='color:#64748B;font-size:15px;'>●</span> <b>Pendente</b> — sem visita registrada<br>
            <span style='color:#DA291C;font-size:15px;'>◆</span> <b>Anomalia crítica</b> detectada
            </div>""",
            unsafe_allow_html=True
        )

    df_mapa = df_filtrado.copy()
    if 'latitude' in df_mapa.columns and 'longitude' in df_mapa.columns:
        df_mapa['latitude']  = pd.to_numeric(df_mapa['latitude'],  errors='coerce')
        df_mapa['longitude'] = pd.to_numeric(df_mapa['longitude'], errors='coerce')
        df_mapa = df_mapa.dropna(subset=['latitude', 'longitude'])
        df_mapa = df_mapa[(df_mapa['latitude'] != 0.0) & (df_mapa['longitude'] != 0.0)].reset_index(drop=True)

        cores_mapa_filtrado = []
        status_lista        = []
        tem_critico_lista   = []

        for _, row_m in df_mapa.iterrows():
            fotos_row = json.loads(row_m['fotos_json'] or '[]')
            tem_critico = any(
                {'Crítico': 'Critico', 'Observação': 'Observacao'}.get(
                    f.get('severidade', 'Normal'), f.get('severidade', 'Normal')
                ) == 'Critico'
                for f in fotos_row
            )
            tem_critico_lista.append(tem_critico)

            if len(fotos_row) > 0:
                status = "Concluída"
                cor    = [218, 41, 28, 220] if tem_critico else [22, 163, 74, 220]
            else:
                status = "Em Andamento"
                cor    = [217, 119, 6, 220]

            status_lista.append(status)
            cores_mapa_filtrado.append(cor)

        df_mapa['status_visita'] = status_lista
        df_mapa['color_rgb']     = cores_mapa_filtrado
        df_mapa['tem_critico']   = tem_critico_lista
        df_mapa['icone_status']  = df_mapa['status_visita'].map({"Concluída": "✅", "Em Andamento": "⚙️", "Pendente": "🕐"})
        df_mapa['alerta_critico'] = df_mapa['tem_critico'].apply(lambda x: "⚠️ ANOMALIA CRÍTICA" if x else "")

        if status_filtro == "✅ Concluída": df_mapa = df_mapa[df_mapa['status_visita'] == "Concluída"]
        elif status_filtro == "⚙️ Em Andamento": df_mapa = df_mapa[df_mapa['status_visita'] == "Em Andamento"]
        elif status_filtro == "🕐 Pendente": df_mapa = df_mapa[df_mapa['status_visita'] == "Pendente"]

        total_df = df_filtrado.copy()
        total_df['fotos_row'] = total_df['fotos_json'].apply(lambda x: json.loads(x or '[]'))
        n_concluidas   = total_df['fotos_row'].apply(lambda f: len(f) > 0).sum()
        n_andamento    = total_df['fotos_row'].apply(lambda f: len(f) == 0).sum()

        mc1, mc2, mc3 = st.columns(3)
        mc1.markdown(f'<div class="eng-metric"><div class="eng-metric-val" style="color:#16A34A;">{n_concluidas}</div><div class="eng-metric-label">CONCLUÍDAS</div></div>', unsafe_allow_html=True)
        mc2.markdown(f'<div class="eng-metric"><div class="eng-metric-val" style="color:#D97706;">{n_andamento}</div><div class="eng-metric-label">EM ANDAMENTO</div></div>', unsafe_allow_html=True)
        mc3.markdown(f'<div class="eng-metric"><div class="eng-metric-val" style="color:#DA291C;">{int(df_filtrado["fotos_json"].apply(lambda x: any({"Crítico":"Critico","Observação":"Observacao"}.get(f.get("severidade","Normal"),f.get("severidade","Normal"))=="Critico" for f in json.loads(x or "[]"))).sum())}</div><div class="eng-metric-label">C/ ANOMALIA CRÍTICA</div></div>', unsafe_allow_html=True)

        st.markdown("<br>", unsafe_allow_html=True)

        if not df_mapa.empty:
            view_state = pdk.ViewState(latitude=df_mapa['latitude'].mean(), longitude=df_mapa['longitude'].mean(), zoom=11, pitch=0)
            layer_pontos = pdk.Layer(
                "ScatterplotLayer", data=df_mapa, get_position="[longitude, latitude]", get_color="color_rgb",
                get_radius=220, radiusMinPixels=10, radiusMaxPixels=20, pickable=True, stroked=True,
                get_line_color=[255, 255, 255, 200], lineWidthMinPixels=2
            )
            tooltip = {
                "html": "<div style='font-family:sans-serif;min-width:200px;'><b style='font-size:13px;'>📍 {site_id}</b><br><span style='font-size:11px;'>{icone_status} <b>{status_visita}</b></span><span style='color:#DA291C;font-weight:bold;'> {alerta_critico}</span><br>👷 {tecnico}<br>📅 {data_hora}<br>🏠 {endereco}</div>",
                "style": {"backgroundColor": "#002060", "color": "white", "borderRadius": "8px", "padding": "12px", "fontSize": "12px"}
            }
            st.pydeck_chart(pdk.Deck(layers=[layer_pontos], initial_view_state=view_state, tooltip=tooltip, map_style="road"), use_container_width=True)
        else:
            st.info(f"💡 Nenhuma visita com status '{status_filtro}' possui coordenadas de GPS cadastradas.")
    else:
        st.info("💡 O banco de dados atual não possui as colunas de Latitude/Longitude preenchidas.")

    st.markdown("---")
    secao("📅", "TENDÊNCIA TEMPORAL DE VISTORIAS")
    df_tempo = df_filtrado.dropna(subset=['data_formatada']).groupby('data_formatada').size().reset_index(name='Laudos')
    if not df_tempo.empty:
        fig_line = px.line(df_tempo, x='data_formatada', y='Laudos', markers=True, labels={"data_formatada": "Data da Vistoria"})
        fig_line.update_traces(line_color=COR_AZUL_MED)
        fig_line.update_layout(margin=dict(l=0, r=0, t=30, b=0))
        st.plotly_chart(fig_line, use_container_width=True)
    else:
        st.info("Dados de data insuficientes para gerar a linha do tempo.")

    st.markdown("---")
    secao("💰", "PAINEL DE ORÇAMENTO — MATERIAIS NECESSÁRIOS")
    itens_orc = []
    for _, row_orc in df_filtrado.iterrows():
        fotos_orc = json.loads(row_orc['fotos_json'] or '[]')
        for foto_orc in fotos_orc:
            mats_orc = foto_orc.get('materiais') or []
            if not mats_orc and foto_orc.get('material_necessario','').strip():
                mats_orc = [{'descricao': foto_orc['material_necessario'], 'unidade':'un', 'quantidade':1, 'custo_unit':0.0}]
            for m in mats_orc:
                if m.get('descricao','').strip():
                    itens_orc.append({
                        'Site':      row_orc['site_id'],
                        'Evidência': foto_orc.get('titulo','—'),
                        'Severidade': foto_orc.get('severidade','Normal'),
                        'Prazo':     foto_orc.get('prazo_correcao','Monitorar'),
                        'Material':  m.get('descricao',''),
                        'Unidade':   m.get('unidade','un'),
                        'Qtd':       float(m.get('quantidade',1)),
                        'Unit_R$':   float(m.get('custo_unit',0.0)),
                        'Total_R$':  float(m.get('quantidade',1)) * float(m.get('custo_unit',0.0)),
                    })

    if not itens_orc:
        st.info("Nenhum material estruturado cadastrado nas evidências dos laudos filtrados.")
    else:
        df_orc = pd.DataFrame(itens_orc)
        total_geral   = df_orc['Total_R$'].sum()
        n_criticos    = df_orc[df_orc['Severidade'].isin(['Critico','Crítico'])].shape[0]
        n_materiais   = df_orc['Material'].nunique()
        total_semcusto = df_orc[df_orc['Unit_R$'] == 0].shape[0]

        oc1, oc2, oc3, oc4 = st.columns(4)
        oc1.markdown(f'<div class="eng-metric"><div class="eng-metric-val" style="color:{COR_AZUL};">R$ {total_geral:,.2f}</div><div class="eng-metric-label">TOTAL ESTIMADO</div></div>', unsafe_allow_html=True)
        oc2.markdown(f'<div class="eng-metric"><div class="eng-metric-val" style="color:#DA291C;">{n_criticos}</div><div class="eng-metric-label">ITENS CRÍTICOS</div></div>', unsafe_allow_html=True)
        oc3.markdown(f'<div class="eng-metric"><div class="eng-metric-val">{n_materiais}</div><div class="eng-metric-label">MATERIAIS DISTINTOS</div></div>', unsafe_allow_html=True)
        oc4.markdown(f'<div class="eng-metric"><div class="eng-metric-val" style="color:#D97706;">{total_semcusto}</div><div class="eng-metric-label">SEM CUSTO INFORMADO</div></div>', unsafe_allow_html=True)

        st.markdown("<br>", unsafe_allow_html=True)
        df_orc_view = df_orc.copy()
        df_orc_view['Qtd'] = df_orc_view['Qtd'].apply(lambda x: f"{x:.0f}")
        df_orc_view['Unit_R$']  = df_orc_view['Unit_R$'].apply(lambda x: f"R$ {x:,.2f}")
        df_orc_view['Total_R$'] = df_orc_view['Total_R$'].apply(lambda x: f"R$ {x:,.2f}")
        st.dataframe(df_orc_view, hide_index=True, use_container_width=True)

    st.markdown("---")
    secao("🛡️", "LOG DE AUDITORIA (LGPD)")
    with st.expander("Visualizar Registros de Sistema"):
        try:
            with sqlite3.connect(DB_NAME) as conn:
                df_audit = pd.read_sql_query("SELECT * FROM audit_log ORDER BY id DESC LIMIT 100", conn)
                if not df_audit.empty:
                    st.dataframe(df_audit, use_container_width=True, hide_index=True)
                else:
                    st.info("Nenhum registro de auditoria encontrado.")
        except:
            st.info("Log de auditoria em inicialização.")

# ══════════════════════════════════════════════════════════════════════════
# MÓDULO: ROTEIRIZAÇÃO TÁTICA (VRP E FIELD SERVICE)
# ══════════════════════════════════════════════════════════════════════════
def calcular_distancia_haversine(lon1, lat1, lon2, lat2):
    R = 6371.0 
    lon1, lat1, lon2, lat2 = map(math.radians, [lon1, lat1, lon2, lat2])
    dlon = lon2 - lon1
    dlat = lat2 - lat1
    a = math.sin(dlat/2)**2 + math.cos(lat1) * math.cos(lat2) * math.sin(dlon/2)**2
    c = 2 * math.atan2(math.sqrt(a), math.sqrt(1-a))
    return R * c

def resolver_tsp_local(ponto_partida, lista_sites):
    rota = [ponto_partida]
    nao_visitados = lista_sites.copy()
    atual = ponto_partida
    
    while nao_visitados:
        proximo = min(nao_visitados, key=lambda x: calcular_distancia_haversine(atual['lon'], atual['lat'], x['lon'], x['lat']))
        rota.append(proximo)
        nao_visitados.remove(proximo)
        atual = proximo
        
    return rota

def formatar_tempo(segundos):
    horas = int(segundos // 3600)
    minutos = int((segundos % 3600) // 60)
    return f"{horas}h {minutos}m"

def tela_roteirizacao():
    banner("ROTEIRIZAÇÃO TÁTICA E FIELD SERVICE")
    
    with sqlite3.connect(DB_NAME) as conn:
        df_sites = pd.read_sql_query("SELECT id, site_id, endereco, latitude, longitude FROM relatorios WHERE latitude != 0.0 AND longitude != 0.0 ORDER BY id DESC", conn)
    
    if df_sites.empty:
        st.info("Nenhum site com coordenadas de GPS cadastrado. Crie laudos com Latitude e Longitude na aba 'Novo Relatório'.")
        return
        
    df_sites = df_sites.drop_duplicates(subset=['site_id']).reset_index(drop=True)
    lista_opcoes = df_sites['site_id'].tolist()
    
    st.markdown("### 📍 Configuração da Rota")
    c_base, c_sites = st.columns(2)
    with c_base:
        site_base = st.selectbox("Ponto de Partida (Base/Hotel):", lista_opcoes)
    with c_sites:
        sites_alvo = st.multiselect("Selecione os Sites a Visitar:", [s for s in lista_opcoes if s != site_base])
        
    ors_token = st.text_input("Token OpenRouteService (Opcional - Deixe em branco para usar TSP Local):", type="password", help="Gere sua chave gratuita em openrouteservice.org para roteamento real nas vias.")
    
    if st.button("🚀 Otimizar Rota de Manutenção", type="primary"):
        if not sites_alvo:
            st.error("Selecione pelo menos 1 site para visitar.")
            return
            
        with st.spinner("Calculando sequenciamento ótimo e projetando métricas de Field Service..."):
            base_row = df_sites[df_sites['site_id'] == site_base].iloc[0]
            pt_partida = {'id': base_row['site_id'], 'lon': float(base_row['longitude']), 'lat': float(base_row['latitude'])}
            
            alvos = []
            for s in sites_alvo:
                row = df_sites[df_sites['site_id'] == s].iloc[0]
                alvos.append({'id': row['site_id'], 'lon': float(row['longitude']), 'lat': float(row['latitude'])})
                
            rota_otimizada = resolver_tsp_local(pt_partida, alvos)
            coords_lista = [[p['lon'], p['lat']] for p in rota_otimizada]
            nomes_rota = [p['id'] for p in rota_otimizada]
            
            distancia_total_km = 0
            duracao_total_seg = 0
            geojson_rota = None
            
            if ors_token.strip():
                try:
                    headers = {
                        'Accept': 'application/json, application/geo+json, application/gpx+xml, img/png; charset=utf-8',
                        'Authorization': ors_token.strip(),
                        'Content-Type': 'application/json; charset=utf-8'
                    }
                    payload = {"coordinates": coords_lista}
                    req = requests.post('https://api.openrouteservice.org/v2/directions/driving-car/geojson', json=payload, headers=headers, timeout=10)
                    
                    if req.status_code == 200:
                        geojson_rota = req.json()
                        distancia_total_km = geojson_rota['features'][0]['properties']['summary']['distance'] / 1000.0
                        duracao_total_seg = geojson_rota['features'][0]['properties']['summary']['duration']
                    else:
                        st.warning(f"Erro na API ORS (Usando cálculo Haversine): {req.text}")
                except Exception as e:
                    st.warning(f"Falha de conexão com a API ORS (Usando cálculo Haversine): {e}")

            if not geojson_rota:
                for i in range(len(rota_otimizada)-1):
                    dist_trecho = calcular_distancia_haversine(rota_otimizada[i]['lon'], rota_otimizada[i]['lat'], rota_otimizada[i+1]['lon'], rota_otimizada[i+1]['lat'])
                    distancia_total_km += dist_trecho
                
                duracao_total_seg = (distancia_total_km / 40.0) * 3600
            
            secao("KPI", "MÉTRICAS DA OPERAÇÃO DE CAMPO")
            m1, m2, m3, m4 = st.columns(4)
            m1.markdown(f'<div class="eng-metric"><div class="eng-metric-val">{len(sites_alvo)}</div><div class="eng-metric-label">SITES ATENDIDOS</div></div>', unsafe_allow_html=True)
            m2.markdown(f'<div class="eng-metric"><div class="eng-metric-val">{distancia_total_km:.1f} km</div><div class="eng-metric-label">QUILOMETRAGEM ESTIMADA</div></div>', unsafe_allow_html=True)
            m3.markdown(f'<div class="eng-metric"><div class="eng-metric-val">{formatar_tempo(duracao_total_seg)}</div><div class="eng-metric-label">WINDSHIELD TIME (DIREÇÃO)</div></div>', unsafe_allow_html=True)
            
            densidade = len(sites_alvo) / distancia_total_km if distancia_total_km > 0 else 0
            m4.markdown(f'<div class="eng-metric"><div class="eng-metric-val">{densidade:.2f}</div><div class="eng-metric-label">DENSIDADE (Sites/Km)</div></div>', unsafe_allow_html=True)

            st.markdown("---")
            seq_html = " &nbsp; ➔ &nbsp; ".join([f"**{n}**" for n in nomes_rota])
            st.info(f"**Ordem Operacional Sugerida:** {seq_html}")

            df_rota = pd.DataFrame(rota_otimizada)
            df_rota['color_rgb'] = [[22, 163, 74, 200] if i == 0 else [0, 48, 135, 200] for i in range(len(df_rota))]
            df_rota['seq_label'] = ["0"] + [str(i) for i in range(1, len(df_rota))]
            df_rota['seq_label'] = df_rota['seq_label'].astype(str)
            
            view_state = pdk.ViewState(latitude=df_rota['lat'].mean(), longitude=df_rota['lon'].mean(), zoom=11, pitch=0)
            
            layers_mapa = [
                pdk.Layer("ScatterplotLayer", data=df_rota, get_position="[lon, lat]", get_color="color_rgb", get_radius=300, radiusMinPixels=15, radiusMaxPixels=25, pickable=True, stroked=True, get_line_color=[255, 255, 255, 200], lineWidthMinPixels=2),
                pdk.Layer("TextLayer", data=df_rota, get_position="[lon, lat]", get_text="seq_label", get_color=[255, 255, 255, 255], get_size=22, sizeScale=1, get_alignment_baseline="'center'", get_text_anchor="'middle'")
            ]
            
            if geojson_rota:
                layer_linha = pdk.Layer("GeoJsonLayer", data=geojson_rota, pickable=False, stroked=True, filled=False, extruded=False, get_line_color=[218, 41, 28, 255], get_line_width=15, lineWidthMinPixels=3)
                layers_mapa.append(layer_linha)
            else:
                path_data = pd.DataFrame([{"path": coords_lista}])
                layer_linha = pdk.Layer("PathLayer", data=path_data, get_path="path", get_color=[218, 41, 28, 200], width_scale=20, width_min_pixels=3, get_width=5)
                layers_mapa.append(layer_linha)

            r = pdk.Deck(layers=layers_mapa, initial_view_state=view_state, tooltip={"text": "Site: {id}"}, map_style="road")
            st.pydeck_chart(r, use_container_width=True)
            st.markdown(f"<span style='color:#16A34A;font-weight:bold;'>🟢 Base/Origem (0)</span> &nbsp;&nbsp; | &nbsp;&nbsp; <span style='color:{COR_AZUL};font-weight:bold;'>🔵 Sites Alvo (Sequência)</span>", unsafe_allow_html=True)

            st.session_state["_rota_resultado"] = {"rota": rota_otimizada, "distancia_km": distancia_total_km, "duracao_seg": duracao_total_seg}
            st.session_state.pop("_rota_pdf_bytes", None)
            st.session_state.pop("_rota_xlsx_bytes", None)

    if st.session_state.get("_rota_resultado"):
        _res = st.session_state["_rota_resultado"]
        rota_salva       = _res["rota"]
        dist_salva       = _res["distancia_km"]
        dur_salva        = _res["duracao_seg"]
        tecnico_rota     = st.session_state.get("_tecnico_global", "N/I")

        st.markdown("---")
        secao("📸", "EVIDÊNCIAS POR SITE — APONTAMENTO DE MATERIAIS")

        evidencias_por_site = {}
        with sqlite3.connect(DB_NAME) as conn_ev:
            for p_ev in rota_salva[1:]:
                sid_ev = p_ev['id']
                row_ev = conn_ev.execute("SELECT fotos_json FROM relatorios WHERE TRIM(UPPER(site_id)) = TRIM(UPPER(?)) ORDER BY id DESC LIMIT 1", (sid_ev,)).fetchone()
                fotos_ev = json.loads(row_ev[0]) if row_ev and row_ev[0] else []
                evidencias_por_site[sid_ev] = fotos_ev

        total_criticos_ui = sum(1 for evs in evidencias_por_site.values() for ev in evs if ev.get('severidade', '') in ('Critico', 'Crítico'))
        total_obs_ui = sum(1 for evs in evidencias_por_site.values() for ev in evs if ev.get('severidade', '') in ('Observacao', 'Observação'))
        total_fotos_ui = sum(len(evs) for evs in evidencias_por_site.values())

        sc1, sc2, sc3 = st.columns(3)
        sc1.markdown(f'<div class="eng-metric"><div class="eng-metric-val" style="color:#DA291C;">{total_criticos_ui}</div><div class="eng-metric-label">EVIDÊNCIAS CRÍTICAS</div></div>', unsafe_allow_html=True)
        sc2.markdown(f'<div class="eng-metric"><div class="eng-metric-val" style="color:#D97706;">{total_obs_ui}</div><div class="eng-metric-label">OBSERVAÇÕES</div></div>', unsafe_allow_html=True)
        sc3.markdown(f'<div class="eng-metric"><div class="eng-metric-val">{total_fotos_ui}</div><div class="eng-metric-label">TOTAL DE EVIDÊNCIAS</div></div>', unsafe_allow_html=True)

        if total_fotos_ui == 0:
            st.info("ℹ️ Nenhuma evidência fotográfica encontrada para os sites selecionados.")
        else:
            st.markdown("**Aponte os materiais necessários para cada evidência antes de exportar:**")

        for seq_idx, p_site in enumerate(rota_salva[1:], 1):
            sid = p_site['id']
            evs = evidencias_por_site.get(sid, [])
            with st.expander(f"📍 [{seq_idx:02d}] {sid} — {len(evs)} evidência(s)", expanded=(len(evs) > 0 and seq_idx == 1)):
                if not evs:
                    st.info("Nenhuma evidência fotográfica registrada para este site.")
                    continue
                for ev_idx, ev in enumerate(evs):
                    sev = ev.get('severidade', 'Normal')
                    cor_sev = COR_VERMELHO if sev in ('Critico','Crítico') else (COR_AMARELO if sev in ('Observacao','Observação') else COR_VERDE)
                    col_img, col_info = st.columns([1, 3])
                    b64_ev = _obter_b64_de_foto(ev)
                    if b64_ev:
                        col_img.markdown(f'<img src="data:image/jpeg;base64,{b64_ev}" style="width:100%;border-radius:6px;border:2px solid {cor_sev};"/>', unsafe_allow_html=True)
                    with col_info:
                        st.markdown(f"**{sanitizar(ev.get('titulo','Evidência'))}** &nbsp; <span style='background:{cor_sev};color:#fff;padding:1px 8px;border-radius:10px;font-size:11px;'>{sanitizar(sev)}</span>", unsafe_allow_html=True)
                        st.caption(sanitizar(ev.get('comentarios', '')))
                        mat_key = f"mat_{sid}_{ev_idx}"
                        mat_val = st.text_input(
                            "🔧 Material necessário para correção:",
                            value=st.session_state.get(mat_key, ev.get('material_necessario', '')),
                            placeholder="Ex: Cabo 6mm², Disjuntor 40A, Conector tipo Y...",
                            key=mat_key
                        )
                        ev['material_necessario'] = mat_val
                    st.markdown("---")

        st.markdown("---")
        secao("📤", "EXPORTAR ROTEIRO")
        col_pdf, col_xlsx = st.columns(2)

        with col_pdf:
            if st.button("📄 Gerar PDF do Roteiro", type="primary", use_container_width=True):
                with st.spinner("Gerando PDF..."):
                    _pb, _pn = gerar_pdf_rota(rota_salva, dist_salva, dur_salva, tecnico_rota, evidencias_por_site)
                st.session_state["_rota_pdf_bytes"] = _pb
                st.session_state["_rota_pdf_nome"]  = _pn
                st.rerun()

            if st.session_state.get("_rota_pdf_bytes"):
                st.download_button(label="⬇️ Baixar PDF", data=st.session_state["_rota_pdf_bytes"], file_name=st.session_state["_rota_pdf_nome"], mime="application/pdf", use_container_width=True, key="dl_pdf_rota")

        with col_xlsx:
            if st.button("📊 Gerar Planilha Excel", use_container_width=True):
                with st.spinner("Gerando Excel..."):
                    _xb, _xn = gerar_excel_rota(rota_salva, dist_salva, dur_salva, tecnico_rota, evidencias_por_site)
                st.session_state["_rota_xlsx_bytes"] = _xb
                st.session_state["_rota_xlsx_nome"]  = _xn
                st.rerun()

            if st.session_state.get("_rota_xlsx_bytes"):
                st.download_button(label="⬇️ Baixar Excel", data=st.session_state["_rota_xlsx_bytes"], file_name=st.session_state["_rota_xlsx_nome"], mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", use_container_width=True, key="dl_xlsx_rota")

def _css_rota_pdf() -> str:
    return f"""
    @page {{ size: A4; margin: 14mm 14mm 18mm 14mm;
        @bottom-left {{ content: "GIRCP — Roteirização Tática • Uso Interno"; font-size: 7pt; color: {COR_CINZA}; font-family: 'Segoe UI', Arial, sans-serif; }}
        @bottom-right {{ content: "Página " counter(page) " de " counter(pages); font-size: 7pt; color: {COR_CINZA}; font-family: 'Segoe UI', Arial, sans-serif; }}
    }}
    * {{ box-sizing: border-box; margin: 0; padding: 0; }}
    body {{ font-family: 'Segoe UI', Helvetica, Arial, sans-serif; color: {COR_TEXTO}; font-size: 9.5pt; line-height: 1.5; background: #fff; }}
    .page-header {{ display: flex; justify-content: space-between; align-items: flex-end; border-bottom: 3px solid {COR_AZUL}; padding-bottom: 10px; margin-bottom: 22px; }}
    .header-titulo {{ color: {COR_AZUL}; font-size: 18pt; font-weight: 900; text-transform: uppercase; letter-spacing: 0.8px; }}
    .header-acento {{ display: inline-block; width: 36px; height: 4px; background: {COR_VERMELHO}; margin-bottom: 4px; }}
    .header-data {{ text-align: right; color: {COR_CINZA}; font-size: 9pt; }}
    .kpi-grid {{ display: flex; gap: 12px; margin-bottom: 22px; }}
    .kpi-box {{ flex: 1; border: 1.5px solid {COR_BORDA}; border-radius: 6px; padding: 12px 10px; text-align: center; background: {COR_AZUL_LIGHT}; }}
    .kpi-val {{ font-size: 18pt; font-weight: 900; color: {COR_AZUL}; }}
    .kpi-lbl {{ font-size: 7pt; color: {COR_CINZA}; text-transform: uppercase; margin-top: 2px; }}
    .section-header {{ background: {COR_AZUL}; color: #fff; font-weight: 700; font-size: 10pt; padding: 7px 14px; margin-top: 22px; margin-bottom: 14px; text-transform: uppercase; border-left: 5px solid {COR_VERMELHO}; border-radius: 2px; }}
    .seq-table {{ width: 100%; border-collapse: collapse; margin-bottom: 18px; }}
    .seq-table th {{ background: {COR_AZUL}; color: #fff; padding: 8px 10px; font-size: 8.5pt; text-align: left; }}
    .seq-table td {{ border: 1px solid {COR_BORDA}; padding: 8px 10px; font-size: 9pt; vertical-align: top; }}
    .seq-table tr:nth-child(even) td {{ background: {COR_AZUL_LIGHT}; }}
    .badge-base {{ background: #16A34A; color: #fff; padding: 1px 8px; border-radius: 10px; font-size: 7pt; font-weight: 700; }}
    .badge-site {{ background: {COR_AZUL}; color: #fff; padding: 1px 8px; border-radius: 10px; font-size: 7pt; font-weight: 700; }}
    .mat-table {{ width: 100%; border-collapse: collapse; margin-bottom: 10px; page-break-inside: avoid; }}
    .mat-table th {{ background: {COR_AZUL_MED}; color: #fff; padding: 6px 10px; font-size: 8pt; text-align: left; }}
    .mat-table td {{ border: 1px solid {COR_BORDA}; padding: 6px 10px; font-size: 8.5pt; vertical-align: top; }}
    .mat-table tr:nth-child(even) td {{ background: {COR_CINZA_LIGHT}; }}
    .badge-critico {{ background: #DA291C; color: #fff; padding: 1px 7px; border-radius: 10px; font-size: 7pt; font-weight: 700; }}
    .badge-obs {{ background: #D97706; color: #fff; padding: 1px 7px; border-radius: 10px; font-size: 7pt; font-weight: 700; }}
    .badge-normal {{ background: #16A34A; color: #fff; padding: 1px 7px; border-radius: 10px; font-size: 7pt; font-weight: 700; }}
    .foto-thumb {{ width: 80px; height: 60px; object-fit: cover; border-radius: 3px; }}
    """

def gerar_pdf_rota(rota_otimizada, distancia_km, duracao_seg, tecnico, evidencias_por_site) -> tuple[bytes, str]:
    densidade = len(rota_otimizada[1:]) / distancia_km if distancia_km > 0 else 0
    kpi_html = f"""
    <div class="kpi-grid">
        <div class="kpi-box"><div class="kpi-val">{len(rota_otimizada)-1}</div><div class="kpi-lbl">Sites Atendidos</div></div>
        <div class="kpi-box"><div class="kpi-val">{distancia_km:.1f} km</div><div class="kpi-lbl">Quilometragem Est.</div></div>
        <div class="kpi-box"><div class="kpi-val">{formatar_tempo(duracao_seg)}</div><div class="kpi-lbl">Windshield Time</div></div>
        <div class="kpi-box"><div class="kpi-val">{densidade:.2f}</div><div class="kpi-lbl">Sites/Km</div></div>
    </div>"""

    seq_rows = ""
    for i, p in enumerate(rota_otimizada):
        badge = '<span class="badge-base">BASE</span>' if i == 0 else f'<span class="badge-site">{i:02d}</span>'
        seq_rows += f"""<tr>
            <td style="text-align:center;">{badge}</td>
            <td><strong>{sanitizar(p['id'])}</strong></td>
            <td style="font-size:8pt;color:{COR_CINZA};">{p.get('lat', ''):.6f}, {p.get('lon', ''):.6f}</td>
        </tr>"""

    mat_html = ""
    total_criticos = 0
    total_obs = 0
    total_normal = 0
    for site_id, evidencias in evidencias_por_site.items():
        if not evidencias:
            continue
        rows_ev = ""
        for ev in evidencias:
            sev = ev.get('severidade', 'Normal')
            cls_b = {'Crítico': 'badge-critico', 'Critico': 'badge-critico', 'Observação': 'badge-obs', 'Observacao': 'badge-obs'}.get(sev, 'badge-normal')
            if sev in ('Critico', 'Crítico'): total_criticos += 1
            elif sev in ('Observacao', 'Observação'): total_obs += 1
            else: total_normal += 1
            b64 = _obter_b64_de_foto(ev)
            thumb = f'<img class="foto-thumb" src="data:image/jpeg;base64,{b64}"/>' if b64 else "—"
            mat_indicado = ev.get('material_necessario', '').strip() or "—"
            rows_ev += f"""<tr>
                <td style="width:90px;">{thumb}</td>
                <td>{sanitizar(ev.get('titulo','—'))}</td>
                <td style="text-align:center;"><span class="{cls_b}">{sanitizar(sev)}</span></td>
                <td style="font-size:8pt;">{sanitizar(ev.get('comentarios','—'))}</td>
                <td style="font-size:8pt;color:{COR_AZUL};font-weight:600;">{sanitizar(mat_indicado)}</td>
            </tr>"""

        mat_html += f"""
        <div class="section-header">📍 {sanitizar(site_id)} — EVIDÊNCIAS E MATERIAIS</div>
        <table class="mat-table">
            <tr>
                <th style="width:90px;">Foto</th><th>Título</th><th>Severidade</th><th>Descrição Técnica</th><th>Material Necessário</th>
            </tr>
            {rows_ev}
        </table>"""

    resumo_sev_html = f"""
    <div class="kpi-grid" style="margin-top:16px;">
        <div class="kpi-box" style="border-color:#DA291C;"><div class="kpi-val" style="color:#DA291C;">{total_criticos}</div><div class="kpi-lbl">Críticos</div></div>
        <div class="kpi-box" style="border-color:#D97706;"><div class="kpi-val" style="color:#D97706;">{total_obs}</div><div class="kpi-lbl">Observações</div></div>
        <div class="kpi-box" style="border-color:#16A34A;"><div class="kpi-val" style="color:#16A34A;">{total_normal}</div><div class="kpi-lbl">Normais</div></div>
    </div>"""

    html = f"""<!DOCTYPE html>
<html lang="pt-BR"><head><meta charset="utf-8"><style>{_css_rota_pdf()}</style></head>
<body>
<div class="page-header">
  <div><div class="header-acento"></div><div class="header-titulo">⚡ Roteiro de Manutenção</div></div>
  <div class="header-data">Emissão: <strong>{datetime.now().strftime('%d/%m/%Y %H:%M')}</strong><br>Técnico: <strong>{sanitizar(tecnico or 'N/I')}</strong></div>
</div>
<div class="section-header">1 &nbsp; KPIs DA OPERAÇÃO</div>
{kpi_html}
{resumo_sev_html}
<div class="section-header">2 &nbsp; SEQUENCIAMENTO OTIMIZADO</div>
<table class="seq-table">
    <tr><th style="width:60px;">Seq.</th><th>Site / Identificação</th><th>Coordenadas</th></tr>
    {seq_rows}
</table>
<div style="page-break-before:always;"></div>
<div class="section-header">3 &nbsp; PAINEL DE EVIDÊNCIAS E MATERIAIS POR SITE</div>
{mat_html if mat_html else '<p style="color:#64748B;padding:10px;">Nenhuma evidência fotográfica encontrada para os sites selecionados.</p>'}
</body></html>"""

    nome = f"Roteiro_{datetime.now().strftime('%Y%m%d%H%M%S')}.pdf"
    pdf_bytes = HTML(string=html).write_pdf()
    return pdf_bytes, nome

def gerar_excel_rota(rota_otimizada, distancia_km, duracao_seg, tecnico, evidencias_por_site) -> tuple[bytes, str]:
    wb = openpyxl.Workbook()
    azul_fill = PatternFill("solid", fgColor="002060")
    azul_med_fill = PatternFill("solid", fgColor="003087")
    cinza_fill = PatternFill("solid", fgColor="EBF0FA")
    vermelho_fill = PatternFill("solid", fgColor="DA291C")
    amarelo_fill = PatternFill("solid", fgColor="D97706")
    verde_fill = PatternFill("solid", fgColor="16A34A")
    branco_font = Font(color="FFFFFF", bold=True)
    azul_font = Font(color="002060", bold=True)
    borda = Border(left=Side(style='thin', color="CBD5E1"), right=Side(style='thin', color="CBD5E1"), top=Side(style='thin', color="CBD5E1"), bottom=Side(style='thin', color="CBD5E1"))

    def _cab(ws, col, row, texto, fill=None, bold=True, center=False):
        cell = ws.cell(row=row, column=col, value=texto)
        if fill: cell.fill = fill
        cell.font = Font(color="FFFFFF" if fill else "002060", bold=bold, size=10 if fill else 9)
        cell.alignment = Alignment(horizontal="center" if center else "left", vertical="center", wrap_text=True)
        cell.border = borda
        return cell

    def _val(ws, col, row, texto, fill=None):
        cell = ws.cell(row=row, column=col, value=texto)
        cell.fill = fill if fill else PatternFill()
        cell.font = Font(size=9)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = borda
        return cell

    ws1 = wb.active
    ws1.title = "KPIs da Operação"
    ws1.column_dimensions['A'].width = 32
    ws1.column_dimensions['B'].width = 20

    ws1.merge_cells("A1:B1")
    c = ws1["A1"]
    c.value = "⚡ GIRCP — ROTEIRO DE MANUTENÇÃO"
    c.fill = azul_fill
    c.font = Font(color="FFFFFF", bold=True, size=14)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 32

    kpis = [
        ("Técnico em Campo", tecnico or "N/I"),
        ("Data/Hora de Emissão", datetime.now().strftime('%d/%m/%Y %H:%M')),
        ("Sites Atendidos", len(rota_otimizada) - 1),
        ("Quilometragem Estimada (km)", f"{distancia_km:.1f}"),
        ("Windshield Time", formatar_tempo(duracao_seg)),
        ("Densidade (Sites/km)", f"{(len(rota_otimizada)-1)/distancia_km:.2f}" if distancia_km > 0 else "—"),
    ]
    for i, (k, v) in enumerate(kpis, start=2):
        ws1.row_dimensions[i].height = 20
        _cab(ws1, 1, i, k)
        _val(ws1, 2, i, v)

    ws2 = wb.create_sheet("Sequenciamento")
    for w, col in zip([8, 22, 18, 18], range(1, 5)):
        ws2.column_dimensions[get_column_letter(col)].width = w
    ws2.row_dimensions[1].height = 22
    for col, hdr in enumerate(["Seq.", "Site / Identificação", "Latitude", "Longitude"], 1):
        _cab(ws2, col, 1, hdr, fill=azul_fill)

    for i, p in enumerate(rota_otimizada):
        r = i + 2
        ws2.row_dimensions[r].height = 18
        fill_r = PatternFill("solid", fgColor="EBF0FA") if i % 2 == 0 else PatternFill()
        label = "BASE" if i == 0 else f"{i:02d}"
        _val(ws2, 1, r, label, fill_r)
        _val(ws2, 2, r, p['id'], fill_r)
        _val(ws2, 3, r, p.get('lat', ''), fill_r)
        _val(ws2, 4, r, p.get('lon', ''), fill_r)

    ws3 = wb.create_sheet("Evidências e Materiais")
    for w, col in zip([22, 26, 16, 36, 30], range(1, 6)):
        ws3.column_dimensions[get_column_letter(col)].width = w
    ws3.row_dimensions[1].height = 22
    for col, hdr in enumerate(["Site", "Título da Evidência", "Severidade", "Descrição Técnica", "Material Necessário"], 1):
        _cab(ws3, col, 1, hdr, fill=azul_fill)

    row_ev = 2
    sev_fills = {'Critico': vermelho_fill, 'Crítico': vermelho_fill, 'Observacao': amarelo_fill, 'Observação': amarelo_fill}
    for site_id, evidencias in evidencias_por_site.items():
        for ev in evidencias:
            sev = ev.get('severidade', 'Normal')
            sfill = sev_fills.get(sev, verde_fill)
            ws3.row_dimensions[row_ev].height = 18
            row_fill = PatternFill("solid", fgColor="EBF0FA") if row_ev % 2 == 0 else PatternFill()
            _val(ws3, 1, row_ev, site_id, row_fill)
            _val(ws3, 2, row_ev, ev.get('titulo', '—'), row_fill)
            c_sev = ws3.cell(row=row_ev, column=3, value=sev)
            c_sev.fill = sfill
            c_sev.font = Font(color="FFFFFF", bold=True, size=9)
            c_sev.alignment = Alignment(horizontal="center", vertical="center")
            c_sev.border = borda
            _val(ws3, 4, row_ev, ev.get('comentarios', '—'), row_fill)
            _val(ws3, 5, row_ev, ev.get('material_necessario', '').strip() or '—', row_fill)
            row_ev += 1

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    nome = f"Roteiro_{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
    return buf.getvalue(), nome

st.set_page_config(page_title="GIRCP | Controle Fotográfico", page_icon="⚡", layout="wide", initial_sidebar_state="expanded")

if not check_password():
    st.stop()

init_db()
aplicar_estilo()

with st.sidebar:
    st.markdown("""
    <div style="text-align:center;padding:20px 0 10px 0;">
      <div style="font-size:42px;">⚡</div>
      <div style="font-size:27px;font-weight:900;letter-spacing:1.5px;color:#fff;">GIRCP</div>
      <div style="font-size:15px;opacity:0.65;color:#fff;margin-top:2px;">SISTEMA CORPORATIVO</div>
    </div>""", unsafe_allow_html=True)
    st.markdown("---")
    tec = st.text_input("👷 TÉCNICO EM CAMPO:", value=st.session_state.get("_tecnico_global", ""), placeholder="Seu nome", key="_tec_sidebar_rel")
    if tec.strip(): st.session_state["_tecnico_global"] = tec.strip()
    st.markdown("---")
    menu = st.radio("NAVEGAÇÃO:", ["📝 NOVO RELATÓRIO", "🔍 PESQUISAR E EXPORTAR", "📊 DASHBOARD", "🗺️ ROTEIRIZAÇÃO TÁTICA"], label_visibility="collapsed")

if menu == "📝 NOVO RELATÓRIO": tela_novo()
elif menu == "🔍 PESQUISAR E EXPORTAR": tela_pesquisa()
elif menu == "📊 DASHBOARD": tela_dashboard()
elif menu == "🗺️ ROTEIRIZAÇÃO TÁTICA": tela_roteirizacao()